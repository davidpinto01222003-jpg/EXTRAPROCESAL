"""
Especializa la búsqueda en Google Drive para el informe
"3. CONTROL PROCESOS EJECUTIVOS ESSA...xlsm" (hoja `HOJA_EXCEL_CONTROL`,
por defecto `ACTIVOS` -- la hoja real que se mantiene al día. NO usa
`DatosProcesados1` ni las demás hojas `DatosProcesadosN`: son copias
aplanadas que no se actualizan solas cuando se edita `ACTIVOS`, así que
pueden quedar con radicados/cuentas/demandados desactualizados o con
valores de relleno. En `ACTIVOS`, un proceso "acumulado" (varias
cuentas bajo un mismo radicado) no repite el No./ESTADO PROCESAL/
RADICADO/JUZGADO en cada fila -- solo la primera fila del grupo los
trae; `leer_procesos_control` los "arrastra hacia abajo" (forward-fill)
a las filas de continuación (identificadas por tener el "No." en
blanco pero sí una CUENTA propia).

A diferencia de `buscar_faltantes_en_drive.py` (que descarga TODO el
contenido relacionado con el radicado), este script trabaja FILA POR
FILA del Excel (no agrupa por número de proceso) y aplica dos reglas
de negocio según el `ESTADO PROCESAL` de cada fila:

1. Cualquier fila que NO sea "terminada" (ver regla 2) -- `ACTIVO`,
   `ACTIVOS CON TITULOS`, `SUSPENDIDO`, `REORGANIZACION`, `REMITIDA A
   CASTIGO`/`A PREPAGO`, etc: la carpeta se nombra `"<numero>.
   <radicado>"` si la fila ya tiene un radicado válido de 23 dígitos,
   o `"<numero>. <ESTADO PROCESAL>"` si todavía no lo tiene (mismo
   formato que la regla 2). SOLO se sube información NO procesal:
   tutelas, derechos de petición, y pagos oficiosos (ver
   `es_informacion_no_procesal`) -- se busca en Google Drive por
   radicado/radicado corto/cuenta (igual que el resto del proyecto), y
   (opcional, ver `BUSCAR_EN_CORREO`) en TODO el Gmail, pero en Gmail
   la búsqueda es AL REVÉS: en vez de buscar por radicado, se busca
   directamente por tutela/derecho de petición/pago oficioso (una sola vez
   para toda la corrida, no una vez por proceso) y CADA correo
   encontrado se empareja después con el proceso correcto si coincide
   su radicado, su cuenta, O el nombre de su demandado (cualquiera de
   los tres alcanza) -- ver `buscar_correo_global_informacion_no_procesal`
   y `_procesos_que_coinciden_con_correo`. Esto encuentra correos que
   una búsqueda por radicado se perdería (ej. un derecho de petición
   que solo menciona el nombre del demandado o la cuenta, no el
   radicado exacto).

2. Procesos terminados por pago, por auto, por contrato/prepago, o que
   nunca se presentaron (`ESTADO PROCESAL` que empieza con `TERMINADO`
   o `NO INICIO` -- igual que `crear_carpetas_terminados_castigo.py`):
   SOLO se sube el documento que deja constancia de que el proceso NO
   sigue su curso -- un auto de terminación, de aceptación de retiro de
   la demanda, de desistimiento, etc (ver `es_auto_terminador`, no
   exige que aparezca literalmente la palabra "AUTO"). Si no se
   encuentra ese documento en Drive, el proceso queda listado en
   `ARCHIVO_PENDIENTES_TERMINADOS` para que lo descargues a mano.

   Dónde queda ese documento depende del estado:
     - `TERMINADO POR AUTO`: NO lleva carpeta propia. El auto se guarda
       RENOMBRADO con el número de proceso del Excel -- `"<numero>.
       <ESTADO PROCESAL EXACTO>"` + su extensión, ej. `"245. TERMINADO
       POR AUTO.pdf"` -- dentro de la carpeta única
       `PROCESOS TERMINADOS POR AUTO` (ver `CARPETA_TERMINADOS_POR_AUTO`),
       para tenerlos todos juntos y no repartidos en cientos de
       carpetas de un solo archivo. Si un proceso trae más de un auto
       (ej. primera y segunda instancia), el segundo queda como
       `"245. TERMINADO POR AUTO_2.pdf"` -- nunca se pisa nada. Los
       autos que corridas ANTERIORES hayan dejado en la carpeta propia
       del proceso se recogen y renombran ahí mismo al empezar cada
       corrida (ver `mover_autos_de_carpetas_antiguas`).
     - Los demás terminados (pago, contrato/prepago, `NO INICIO`):
       como siempre, carpeta propia `"<numero>. <ESTADO PROCESAL
       EXACTO del Excel>"` y el documento conserva el nombre que traía
       en Drive.

IMPORTANTE -- procesos "acumulados": el Excel repite el mismo número de
proceso en más de una fila en dos casos distintos:
  - Cuentas/demandados distintos bajo el MISMO radicado (proceso
    "acumulado": un solo expediente judicial que agrupa varias cuentas).
    Estas filas se FUSIONAN en UNA SOLA carpeta (una carpeta por
    radicado, no una por cuenta) -- sus cuentas y demandados se juntan
    para la búsqueda/validación, pero la carpeta es una sola.
  - Menos frecuente: el mismo número de proceso con un radicado
    DISTINTO en cada fila (numeración administrativa repetida por
    error o por reuso, no es el mismo expediente) -- estas SÍ generan
    carpetas separadas, porque el nombre (que incluye el radicado) ya
    sale distinto para cada una. Ver `clasificar_procesos`.

La clasificación de "información no procesal" y de "documento que
termina el proceso" es por PALABRAS CLAVE (nombre/asunto y, si es
PDF/DOCX o el cuerpo de un correo, su contenido) -- es una heurística,
no perfecta. Cada decisión queda registrada en el log para que la
revises y ajustes las listas de palabras clave
(`PALABRAS_TIPO_INFORMACION_FUERTES`, `PALABRAS_PAGO_OFICIOSO`,
`PALABRAS_PROCESO_NO_CONTINUA`) si hace falta.

Reutiliza toda la infraestructura de `buscar_faltantes_en_drive.py`
(autenticación de Drive, búsqueda por radicado/radicado corto/cuenta,
validación de que el documento sea de ESSA y del demandado correcto,
descarga de PDF/exportables de Google) -- no repite esa lógica. Para
Gmail, reutiliza el mismo mecanismo de búsqueda que ese script (IMAP
`X-GM-RAW` sobre "Todos los mensajes", igual que buscar en la barra de
búsqueda de Gmail), pero además revisa el CUERPO del correo y CUALQUIER
adjunto (no solo enlaces de Drive o adjuntos .zip).

Requiere las mismas credenciales que `buscar_faltantes_en_drive.py`:
`credenciales_drive.json`/`token_drive.json` para Drive (ver README), y
`credenciales_sgde.txt` para Gmail (opcional -- si no existe, la
búsqueda en correo simplemente se omite, ver `BUSCAR_EN_CORREO`).

Respeta MODO_PRUEBA (por defecto True): en modo prueba solo BUSCA y
CLASIFICA, mostrando qué subiría y a qué carpeta, sin crear carpetas ni
descargar nada todavía.

Rendimiento: la mayor parte del tiempo se va esperando la red (Google
Drive), no la CPU, así que este script busca varios procesos EN
PARALELO (ver NUM_HILOS, por defecto 8 a la vez) en vez de uno por uno.
Además, cuando el NOMBRE de un archivo ya alcanza para descartarlo (ej.
"DEMANDA EJECUTIVA.pdf"), ni siquiera se descarga ni se lee su contenido
-- ver `_decision_solo_por_nombre`.
"""

import csv
import email
import imaplib
import io
import logging
import os
import re
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl

import buscar_faltantes_en_drive as buscador
import crear_carpetas_terminados_castigo as terminados_folder
import procesos_juridicos as organizador
import validar_renombrar_carpetas as cruce_excel

# ============================= CONFIGURACION =============================

# Ruta al informe de Excel de procesos ejecutivos (.xlsm).
RUTA_EXCEL_CONTROL = r"C:\Users\Francy\OneDrive\INFORME ENTREGA ESSA\3. CONTROL PROCESOS EJECUTIVOS ESSA 29072026 .xlsm"

# Hoja del Excel a leer -- "ACTIVOS" es la hoja que se mantiene al dia
# (la que editas tu). NO usar "DatosProcesados1" ni las demas hojas
# "DatosProcesadosN": son copias aplanadas que NO se actualizan solas
# cuando editas "ACTIVOS" -- usarlas hacia que el script leyera
# radicados/cuentas/demandados desactualizados o de relleno.
HOJA_EXCEL_CONTROL = "ACTIVOS"

# Fila donde estan los encabezados de columna reales dentro de "ACTIVOS"
# (las primeras filas son titulos/logos de la hoja).
FILA_ENCABEZADO_CONTROL = 5

COLUMNA_NO = "No."
COLUMNA_ESTADO = "ESTADO PROCESAL"
COLUMNA_CUENTA = "CUENTA"
COLUMNA_DEMANDADO = "DEMANDADO"
COLUMNA_JUZGADO = "JUZGADO"
COLUMNA_RADICADO = "RADICADO"

# Prefijos de ESTADO PROCESAL que cuentan como "terminado" para este
# script -- igual que crear_carpetas_terminados_castigo.py: TERMINADO
# (por pago, por auto, por contrato/prepago) y NO INICIO (nunca se
# presentó la demanda, quedó como reclamación administrativa). CUALQUIER
# otro estado (ACTIVO, SUSPENDIDO, REORGANIZACION, REMITIDA A CASTIGO/
# PREPAGO, etc) se organiza con la carpeta "numero. radicado" de siempre.
PREFIJOS_ESTADO_TERMINADO = ("TERMINADO", "NO INICIO")

# Los procesos cuyo ESTADO PROCESAL empieza con "TERMINADO POR AUTO" NO
# llevan una carpeta propia: el auto que los termina se guarda RENOMBRADO
# con el numero de proceso del Excel (ej. "245. TERMINADO POR AUTO.pdf")
# dentro de esta UNICA carpeta comun, para tenerlos todos juntos en vez
# de repartidos en cientos de carpetas de un solo archivo. El resto de
# terminados (por pago, por contrato/prepago, NO INICIO) sigue con su
# carpeta "<numero>. <ESTADO PROCESAL>" de siempre.
PREFIJO_ESTADO_TERMINADO_POR_AUTO = "TERMINADO POR AUTO"
CARPETA_TERMINADOS_POR_AUTO = "PROCESOS TERMINADOS POR AUTO"

# En CARPETA_TERMINADOS_POR_AUTO va UNICAMENTE el auto de cada proceso
# (un archivo por proceso, con el nombre del proceso). Si a un proceso
# ya terminado por auto igual le llega un correo extraprocesal (ver
# revisar_correo_pro.py), ese correo NO se mezcla ahi: se guarda en
# "<numero>. TERMINADO POR AUTO\CORREOS", aparte, para no ensuciar la
# carpeta de autos. Ver carpeta_de_correos().
SUBCARPETA_CORREOS_TERMINADOS_POR_AUTO = "CORREOS"

# Carpeta donde se crean las carpetas de cada proceso. A diferencia del
# resto del proyecto (que usa CARPETA_PROCESOS de validar_renombrar_carpetas.py,
# un disco duro externo detectado por etiqueta), este script tiene su
# PROPIO destino fijo -- no depende de ningun disco externo.
CARPETA_PROCESOS = r"C:\Users\Francy\Documents\INFORMACIÓN EXTRAPROCESAL"

ARCHIVO_LOG = os.path.join(os.path.dirname(__file__), "clasificar_procesos_ejecutivos.log")

# Procesos terminados (por pago/auto/contrato/no inicio) en los que NO
# se encontró el documento que los termina -- para que los descargues a mano.
ARCHIVO_PENDIENTES_TERMINADOS = os.path.join(os.path.dirname(__file__), "terminados_sin_auto_pendientes.csv")

# Carpetas (una por línea) que YA se revisaron por completo en una
# corrida anterior -- en la siguiente corrida se OMITEN por completo
# (ni Drive ni Gmail), para no repetir horas de trabajo ya hecho. Un
# proceso "con radicado" se marca aquí apenas termina su búsqueda en
# Drive (haya encontrado algo o no); uno "terminado" solo se marca
# cuando SÍ se encontró el documento que lo termina -- si quedó
# pendiente, se vuelve a intentar en la próxima corrida. No se marca
# nada mientras MODO_PRUEBA esté activo (una simulación no debe hacer
# que la próxima corrida real se salte procesos sin haberlos tocado de
# verdad). Para forzar que se revise todo de nuevo, borra este archivo.
ARCHIVO_PROCESADOS = os.path.join(os.path.dirname(__file__), "clasificar_procesos_ejecutivos_revisados.txt")

# True (por defecto): no crea carpetas ni descarga nada, solo busca y
# muestra qué haría. False: aplica los cambios de verdad.
MODO_PRUEBA = True

# True (por defecto): ademas de Drive, busca en TODO tu Gmail (via IMAP,
# usando credenciales_sgde.txt -- ver README) tutelas, derechos de
# peticion y pagos oficiosos relacionados con cada proceso. Si ese
# archivo de credenciales no existe, esta busqueda se omite sola, sin
# error.
BUSCAR_EN_CORREO = True

# Cuantos procesos se buscan en Google Drive AL MISMO TIEMPO (en hilos
# separados). Casi todo el tiempo de este script se va esperando la
# respuesta de la API de Drive (red), no procesando en la CPU -- por
# eso buscar varios procesos a la vez (en vez de uno por uno) reduce el
# tiempo total casi en la misma proporcion que NUM_HILOS, sin sobrecargar
# tu computador. 8 es un valor conservador frente al limite de la API de
# Drive (miles de solicitudes por minuto); puedes subirlo (ej. 15-20) si
# tu internet aguanta y quieres que vaya mas rapido todavia, o bajarlo a
# 1 para volver al comportamiento anterior (uno por uno, mas facil de
# leer en el log).
NUM_HILOS = 8

# --------------------- Palabras clave de clasificación ---------------------

# Frases que, si aparecen en el contenido (nombre de archivo, texto del
# documento, o asunto/cuerpo de un correo), confirman que es un derecho
# de petición o una tutela. A propósito NO se incluye "SOLICITUD": es
# una palabra demasiado común dentro de cualquier memorial procesal
# ("se solicita...") y generaba muchas imprecisiones.
PALABRAS_TIPO_INFORMACION_FUERTES = [
    # "PETICION" a secas cubre, de una sola vez, todas las formas en
    # que aparece: "DERECHO DE PETICION", "DERECHOS DE PETICION",
    # "PETICIONES", y -- lo importante -- las RESPUESTAS ("RESPUESTA AL
    # DERECHO DE PETICION", "EN RESPUESTA A SU PETICION", "RTA PETICION"),
    # que tambien son informacion extraprocesal y hay que guardar.
    # Los dos primeros se dejan escritos aunque queden cubiertos por
    # "PETICION", para que se lea de un vistazo que es lo que se busca.
    "DERECHO DE PETICION",
    "DERECHOS DE PETICION",
    "PETICION",
    "ACCION DE TUTELA",
    "TUTELA",
    # PQR / PQRS / PQRSD -- como se rotulan en Colombia las peticiones
    # y, sobre todo, LAS RESPUESTAS a un derecho de peticion ("Respuesta
    # PQRSD No. 12345"). "PQR" alcanza para las tres formas.
    "PQR",
]

# Frases que identifican un pago oficioso -- lo ÚNICO más, aparte de
# petición/tutela, que cuenta como "información no procesal".
PALABRAS_PAGO_OFICIOSO = [
    "PAGO OFICIOSO",
    "PAGOS OFICIOSOS",
    "PAGO DE MANERA OFICIOSA",
]

# Si el NOMBRE del archivo/asunto trae cualquiera de estas frases, es
# un documento procesal (demanda, memorial, solicitud dirigida al
# juzgado, mandamiento, etc) -- se descarta como información no
# procesal aunque en algún lado de su contenido mencione de pasada una
# tutela/petición (ej. narrando el historial del proceso). El nombre
# del archivo es una señal mucho más confiable que rastrear la palabra
# en las paginas de un PDF largo -- ver es_informacion_no_procesal.
#
# A propósito NO se incluye "CONTESTACION" sola: "contestación de la
# demanda" (que sí debe excluirse) ya queda cubierta por "DEMANDA"
# arriba, pero "contestación"/"respuesta" TAMBIÉN es como se suele
# llamar la respuesta que un municipio/departamento envía a un derecho
# de petición o una tutela -- esas respuestas SÍ son información no
# procesal y deben poder pasar a la revisión de contenido (tier 3), no
# descartarse solo por decir "contestación".
PALABRAS_PROCESAL_EXCLUIR = [
    "DEMANDA", "MEMORIAL", "MANDAMIENTO",
    "RECURSO DE REPOSICION", "RECURSO DE APELACION", "EXCEPCIONES",
    "TRASLADO", "SENTENCIA", "LIQUIDACION DE CREDITO",
    "SOLICITUD DE CONCILIACION", "SOLICITA REQUERIR", "SOLICITUD DE REQUERIR",
    "REQUERIMIENTO", "MEDIDA CAUTELAR", "EMBARGO", "SECUESTRO", "PODER",
]

# Frases que dan a entender que el proceso judicial NO sigue su curso
# (terminó por pago, se aceptó el retiro de la demanda, se decretó
# desistimiento o archivo, etc) -- NO hace falta que el documento diga
# literalmente "AUTO", con que aparezca cualquiera de estas frases
# alcanza (ver es_auto_terminador).
PALABRAS_PROCESO_NO_CONTINUA = [
    "TERMINA EL PROCESO", "TERMINACION DEL PROCESO", "SE DA POR TERMINADO",
    "TERMINADO EL PROCESO", "PROCESO TERMINADO", "DECLARA TERMINADO EL PROCESO",
    "TERMINACION POR PAGO", "TERMINACION POR CONTRATO", "TERMINACION DE LA OBLIGACION",
    "DECRETA LA TERMINACION", "DECRETA TERMINACION",
    "ACEPTA EL RETIRO", "ACEPTA RETIRO DE LA DEMANDA", "RETIRO DE LA DEMANDA",
    "SE RETIRA LA DEMANDA", "DESISTIMIENTO DEL PROCESO", "DESISTIMIENTO DE LA DEMANDA",
    "APRUEBA EL DESISTIMIENTO", "ARCHIVA EL PROCESO", "ARCHIVESE EL PROCESO",
    "ARCHIVO DEL PROCESO", "CULMINA EL PROCESO", "FINALIZA EL PROCESO",
    "NO CONTINUA EL PROCESO", "CESE DE LA ACCION EJECUTIVA", "DA POR CONCLUIDO EL PROCESO",
]

# ===========================================================================


def configurar_logging():
    # Con hora (%(asctime)s): para poder ver a simple vista cuanto tarda
    # entre una fila y otra -- por ejemplo entre las filas de un mismo
    # proceso acumulado, para confirmar que la cache de _buscar_archivos/
    # _info_documento de verdad evita repetir trabajo.
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(message)s",
        datefmt="%H:%M:%S",
        handlers=[
            logging.FileHandler(ARCHIVO_LOG, encoding="utf-8", mode="w"),
            logging.StreamHandler(),
        ],
    )


def _normalizar_radicado(valor):
    if valor is None:
        return None
    texto = re.sub(r"[\s\-]", "", str(valor).strip())
    return texto if texto.isdigit() and len(texto) == 23 else None


def _encontrar_columna(encabezados, nombre_buscado):
    for idx, valor in enumerate(encabezados):
        if valor and str(valor).strip().upper() == nombre_buscado.strip().upper():
            return idx
    raise ValueError(
        f"No se encontro la columna '{nombre_buscado}' en la fila {FILA_ENCABEZADO_CONTROL} "
        f"de la hoja '{HOJA_EXCEL_CONTROL}'."
    )


def leer_procesos_control():
    """
    Lee HOJA_EXCEL_CONTROL ("ACTIVOS") fila por fila. En esa hoja, un
    proceso "acumulado" (varias cuentas bajo el mismo radicado) NO
    repite el No./ESTADO PROCESAL/RADICADO/JUZGADO en cada fila -- solo
    la PRIMERA fila del grupo los trae, y las siguientes (con el "No."
    en blanco) solo traen su propia CUENTA/DEMANDADO. Por eso aca se
    "arrastra hacia abajo" (forward-fill) el No./estado/radicado/juzgado
    de la ultima fila con "No." real hacia sus filas de continuacion.
    Una fila con "No." en blanco Y sin CUENTA se ignora (nota, leyenda,
    o fila vacia -- no es una cuenta real de ningun proceso).
    Devuelve una lista de {fila_excel, numero, estado, radicado, cuenta,
    demandado, juzgado} -- una entrada POR FILA (incluidas las de
    continuacion), igual que antes.
    """
    wb = openpyxl.load_workbook(RUTA_EXCEL_CONTROL, data_only=True, read_only=True)
    if HOJA_EXCEL_CONTROL not in wb.sheetnames:
        raise ValueError(
            f"La hoja '{HOJA_EXCEL_CONTROL}' no existe en {RUTA_EXCEL_CONTROL}. "
            f"Hojas disponibles: {wb.sheetnames}"
        )
    ws = wb[HOJA_EXCEL_CONTROL]
    filas = ws.iter_rows(min_row=FILA_ENCABEZADO_CONTROL, max_row=FILA_ENCABEZADO_CONTROL, values_only=True)
    encabezados = next(filas)

    idx_no = _encontrar_columna(encabezados, COLUMNA_NO)
    idx_estado = _encontrar_columna(encabezados, COLUMNA_ESTADO)
    idx_cuenta = _encontrar_columna(encabezados, COLUMNA_CUENTA)
    idx_demandado = _encontrar_columna(encabezados, COLUMNA_DEMANDADO)
    idx_juzgado = _encontrar_columna(encabezados, COLUMNA_JUZGADO)
    idx_radicado = _encontrar_columna(encabezados, COLUMNA_RADICADO)

    procesos = []
    numero_actual = None
    estado_actual = ""
    radicado_actual = None
    juzgado_actual = ""

    for fila_excel, fila in enumerate(
        ws.iter_rows(min_row=FILA_ENCABEZADO_CONTROL + 1, values_only=True),
        start=FILA_ENCABEZADO_CONTROL + 1,
    ):
        numero_crudo = fila[idx_no]
        cuenta_crudo = fila[idx_cuenta]

        if numero_crudo is not None and isinstance(numero_crudo, (int, float)):
            numero_actual = int(numero_crudo)
            estado_actual = str(fila[idx_estado]).strip() if fila[idx_estado] is not None else ""
            radicado_actual = _normalizar_radicado(fila[idx_radicado])
            juzgado_actual = str(fila[idx_juzgado]).strip() if fila[idx_juzgado] is not None else ""
        elif numero_actual is None or cuenta_crudo is None:
            continue  # fila vacia, de notas/leyenda, o sin ningun proceso todavia

        cuenta = str(cuenta_crudo).strip() if cuenta_crudo is not None else ""
        demandado_crudo = fila[idx_demandado]
        demandado = str(demandado_crudo).strip() if demandado_crudo is not None else ""

        procesos.append({
            "fila_excel": fila_excel,
            "numero": numero_actual,
            "estado": estado_actual,
            "radicado": radicado_actual,
            "cuenta": cuenta if cuenta.strip("0") else "",
            "demandado": demandado,
            "juzgado": juzgado_actual,
        })
    return procesos


def es_terminado_por_auto(estado: str) -> bool:
    """
    True si el ESTADO PROCESAL del Excel es "TERMINADO POR AUTO" (o
    empieza por ahi, ej. "TERMINADO POR AUTO QUE ACEPTA EL RETIRO DE LA
    DEMANDA"). Estos procesos no llevan carpeta propia: su auto se
    guarda renombrado dentro de CARPETA_TERMINADOS_POR_AUTO.
    """
    return estado.strip().upper().startswith(PREFIJO_ESTADO_TERMINADO_POR_AUTO)


def carpeta_de_correos(proceso) -> Path:
    """
    Carpeta donde va un CORREO de este proceso (la usa tambien
    revisar_correo_pro.py). Para casi todos es la carpeta del proceso;
    para los "TERMINADO POR AUTO" -- que ya no tienen carpeta propia,
    sino un solo archivo dentro de CARPETA_TERMINADOS_POR_AUTO -- es
    "<numero>. TERMINADO POR AUTO\\CORREOS", para que la carpeta de
    autos quede con un unico archivo por proceso y nada mas.
    """
    if proceso.get("nombre_archivo"):
        return Path(CARPETA_PROCESOS) / proceso["nombre_carpeta"] / SUBCARPETA_CORREOS_TERMINADOS_POR_AUTO
    return Path(CARPETA_PROCESOS) / proceso["nombre_carpeta"]


def clasificar_procesos(procesos):
    """
    Separa las filas leidas en dos listas de trabajo: con_radicado
    (todo lo que no es "terminado" -- se busca información no procesal)
    y terminados (se busca el documento que termina el proceso), mas
    una lista informativa sin_estado. Toda fila con ESTADO PROCESAL
    diligenciado recibe una carpeta: con su radicado si ya lo tiene, o
    con "numero. ESTADO PROCESAL" si todavia no (igual que las filas
    terminadas) -- ver módulo docstring.

    UNA SOLA carpeta por nombre: las filas que producen el MISMO nombre
    (ej. un proceso "acumulado" -- mismo número y mismo radicado en
    varias cuentas) se FUSIONAN en un único proceso de trabajo, sumando
    sus cuentas/demandados/filas del Excel -- nunca se crean carpetas
    "_2", "_3" por tener más de una cuenta.

    EXCEPCIÓN -- "TERMINADO POR AUTO" (ver
    PREFIJO_ESTADO_TERMINADO_POR_AUTO): esos procesos no llevan carpeta
    propia. Su "nombre_carpeta" se sigue calculando igual (es la
    identidad del proceso: la clave para agrupar filas del mismo caso y
    para ARCHIVO_PROCESADOS), pero su destino en el disco es la carpeta
    común CARPETA_TERMINADOS_POR_AUTO, y el auto que se descargue se
    guarda ahí con ese mismo nombre ("nombre_archivo", ej. "245.
    TERMINADO POR AUTO" + la extensión del documento). Por eso cada
    proceso lleva:
      - nombre_carpeta:  identidad del proceso (no cambia).
      - carpeta_destino: carpeta REAL en el disco donde va su contenido.
      - nombre_archivo:  con qué nombre se guarda lo que se descargue,
                         o None para dejarle su nombre original.
    """
    con_radicado_por_nombre = {}
    terminados_por_nombre = {}
    sin_estado = []

    for fila in procesos:
        if not fila["estado"]:
            sin_estado.append(fila)
            continue

        numero, estado = fila["numero"], fila["estado"]

        if estado.upper().startswith(PREFIJOS_ESTADO_TERMINADO):
            base = terminados_folder._nombre_carpeta_para(numero, estado) or f"{numero}. {estado}"
            grupo = terminados_por_nombre
        else:
            base = f"{numero}. {fila['radicado']}" if fila["radicado"] else f"{numero}. {estado}"
            grupo = con_radicado_por_nombre

        nombre_carpeta = organizador.sanear_nombre(base)
        proceso = grupo.get(nombre_carpeta)
        if proceso is None:
            proceso = dict(fila)
            proceso["nombre_carpeta"] = nombre_carpeta
            if es_terminado_por_auto(estado):
                proceso["carpeta_destino"] = CARPETA_TERMINADOS_POR_AUTO
                proceso["nombre_archivo"] = nombre_carpeta
            else:
                proceso["carpeta_destino"] = nombre_carpeta
                proceso["nombre_archivo"] = None
            proceso["cuentas"] = []
            proceso["demandados"] = []
            proceso["filas_excel"] = []
            grupo[nombre_carpeta] = proceso

        if fila["cuenta"] and fila["cuenta"] not in proceso["cuentas"]:
            proceso["cuentas"].append(fila["cuenta"])
        if fila["demandado"] and fila["demandado"] not in proceso["demandados"]:
            proceso["demandados"].append(fila["demandado"])
        proceso["filas_excel"].append(fila["fila_excel"])

    return list(con_radicado_por_nombre.values()), list(terminados_por_nombre.values()), sin_estado


# ==================== Google Drive (usa buscar_faltantes_en_drive.py) ====================


def autenticar_drive_o_none():
    try:
        return buscador.autenticar_drive()
    except Exception as error:
        logging.error("[Drive] No se pudo conectar con Google Drive: %s", error)
        return None


# El cliente de Google Drive (via httplib2) NO es seguro para compartir
# entre varios hilos al mismo tiempo -- cada hilo necesita su PROPIO
# objeto de servicio. El hilo principal ya se autentico antes de lanzar
# los hilos (asi el token ya quedo valido/renovado en disco, ver
# procesar()), asi que cada hilo solo necesita leerlo de nuevo: no
# dispara otro refresh ni el flujo de login.
_hilo_local = threading.local()


def _servicio_del_hilo():
    if not hasattr(_hilo_local, "servicio"):
        _hilo_local.servicio = buscador.autenticar_drive()
    return _hilo_local.servicio


def _terminos_busqueda(radicado, cuentas):
    terminos = []
    if radicado:
        terminos.append(radicado)
        terminos.extend(buscador.radicados_cortos(radicado))
    for cuenta in cuentas:
        if buscador._cuenta_es_valida_para_buscar(cuenta):
            terminos.append(cuenta)
    return terminos


def _candidatos_en_drive(servicio, terminos):
    vistos = {}
    for termino in terminos:
        for item in buscador.buscar_en_drive(servicio, termino):
            vistos[item["id"]] = item
    return list(vistos.values())


# Cache de [(archivo, carpeta), ...] ya encontrados en Drive para un
# radicado, indexado por radicado. Un proceso "acumulado" (mismo numero
# y mismo radicado en varias cuentas -- ver modulo docstring) genera
# una fila por cuenta, y las 7, 10 o mas filas de un mismo acumulado
# comparten la MISMA carpeta de Drive: sin esta cache, cada fila
# repetia la busqueda y volvia a listar/leer los mismos archivos desde
# cero. Solo se cachea cuando hay radicado -- sin el (ej. NO INICIO),
# cada fila depende de su propia cuenta y no hay nada compartido que
# cachear.
_CACHE_ARCHIVOS_POR_RADICADO = {}


def _buscar_archivos(servicio, radicado, cuentas):
    if radicado and radicado in _CACHE_ARCHIVOS_POR_RADICADO:
        return _CACHE_ARCHIVOS_POR_RADICADO[radicado]

    terminos = _terminos_busqueda(radicado, cuentas)
    candidatos = _candidatos_en_drive(servicio, terminos)
    archivos = _archivos_de_candidatos(servicio, candidatos, radicado)

    if radicado:
        _CACHE_ARCHIVOS_POR_RADICADO[radicado] = archivos
    return archivos


def _archivos_de_candidatos(servicio, candidatos, radicado):
    """
    A partir de los candidatos de Drive (carpetas o archivos sueltos que
    coincidieron con el radicado/cuenta), devuelve [(archivo, carpeta), ...]
    de los PDF/exportables relevantes -- reutilizando el mismo criterio
    de "carpeta propia del caso" que buscar_faltantes_en_drive.py.
    """
    vistos = {}
    for item in candidatos:
        carpeta = buscador.carpeta_contenedora(servicio, item)
        if not carpeta:
            continue

        es_propia = item.get("mimeType") == buscador.MIME_CARPETA or buscador._carpeta_es_dedicada_al_caso(
            item, carpeta, radicado or ""
        )
        if es_propia:
            if radicado:
                archivos = buscador._archivos_relacionados(servicio, carpeta, radicado)
            else:
                # Sin radicado real (ej. NO INICIO): no hay como filtrar
                # por nombre, se listan los PDF de primer nivel de la
                # carpeta tal cual -- la validacion de ESSA/demandado
                # (ver _archivo_es_seguro) sigue aplicando igual.
                try:
                    respuesta = servicio.files().list(
                        q=f"'{carpeta['id']}' in parents and trashed = false",
                        fields="files(id, name, mimeType)",
                    ).execute()
                except buscador.HttpError:
                    respuesta = {}
                archivos = [
                    a for a in respuesta.get("files", [])
                    if a.get("mimeType") != buscador.MIME_CARPETA and buscador._es_pdf_o_exportable(a)
                ]
        elif item.get("mimeType") != buscador.MIME_CARPETA and buscador._es_pdf_o_exportable(item):
            archivos = [item]
        else:
            archivos = []

        for archivo in archivos:
            vistos[archivo["id"]] = (archivo, carpeta)
    return list(vistos.values())


# Cache del contenido normalizado (nombre + texto de PDF/DOCX) de cada
# archivo de Drive ya leido en esta corrida -- indexado por el ID del
# archivo. Sin esto, un mismo archivo se terminaba descargando/leyendo
# DOS veces (una para _archivo_es_seguro, otra para clasificarlo), y
# ademas una vez POR CADA fila de un proceso "acumulado" que comparte
# el mismo radicado (ver _buscar_archivos) -- con carpetas de cientos
# de PDF y procesos con varias cuentas, eso se sentia como si el script
# estuviera pegado.
_CACHE_CONTENIDO_ARCHIVO = {}


def _info_documento(servicio, archivo):
    """
    (nombre_normalizado, contenido_normalizado) -- nombre_normalizado es
    SOLO el nombre del archivo (señal mas confiable para clasificar el
    TIPO de documento); contenido_normalizado es nombre + texto de PDF/
    DOCX si aplica (para buscar el tipo tambien dentro del documento).
    Cacheado por ID de archivo.
    """
    archivo_id = archivo.get("id")
    if archivo_id in _CACHE_CONTENIDO_ARCHIVO:
        return _CACHE_CONTENIDO_ARCHIVO[archivo_id]

    nombre = archivo.get("name", "")
    texto = ""
    if Path(nombre).suffix.lower() in buscador.EXTENSIONES_CONTENIDO_DRIVE:
        texto = buscador._texto_de_archivo_drive(servicio, archivo)
    nombre_norm = buscador._normalizar_para_comparar(nombre)
    contenido = buscador._normalizar_para_comparar(nombre + " " + texto)
    resultado = (nombre_norm, contenido)

    if archivo_id:
        _CACHE_CONTENIDO_ARCHIVO[archivo_id] = resultado
    return resultado


def _archivo_es_seguro(nombre_carpeta, archivo, contenido_normalizado, demandados):
    """
    Aplica la regla obligatoria de buscar_faltantes_en_drive.py: el
    documento (o su carpeta contenedora) tiene que mencionar a ESSA/
    Electrificadora de Santander, y no debe contradecir al DEMANDADO
    esperado del Excel. Reutiliza 'contenido_normalizado' (ya leido por
    _info_documento) en vez de volver a descargar el archivo para
    revisar el demandante, como hacia la version anterior.
    """
    contexto_nombre = buscador._normalizar_para_comparar(f"{nombre_carpeta} {archivo.get('name', '')}")
    terminos_essa = [buscador._normalizar_para_comparar(t) for t in buscador.TERMINOS_DEMANDANTE_VALIDO]
    tiene_essa = any(t in contexto_nombre for t in terminos_essa) or any(t in contenido_normalizado for t in terminos_essa)
    if not tiene_essa:
        return False, "no se confirmo que el proceso sea de ESSA/Electrificadora de Santander"

    for demandado in demandados:
        resultado = buscador._demandado_coincide_en_varios([contexto_nombre, contenido_normalizado], demandado)
        if resultado is False:
            return False, f"el nombre en el documento no corresponde al demandado esperado ({demandado})"

    return True, "ok"


def _decision_solo_por_nombre(nombre_normalizado):
    """
    Intenta decidir es_informacion_no_procesal SOLO con el nombre del
    archivo/asunto, SIN mirar el contenido -- devuelve (True, motivo) o
    (False, motivo) si el nombre ya alcanza para decidir (tiers 1 y 2),
    o None si hace falta revisar el contenido (tier 3). Se usa para
    evitar DESCARGAR y leer el contenido de archivos que el nombre ya
    descarta (ej. "DEMANDA EJECUTIVA.pdf") -- con carpetas de decenas de
    documentos por proceso, esto ahorra la mayor parte del trabajo
    pesado (descargar + extraer texto de PDF/DOCX), que es lo que mas
    tiempo toma.
    """
    palabras_objetivo = PALABRAS_TIPO_INFORMACION_FUERTES + PALABRAS_PAGO_OFICIOSO

    if any(frase in nombre_normalizado for frase in palabras_objetivo):
        return True, "tutela/derecho de peticion/pago oficioso (nombre del archivo)"

    if any(marca in nombre_normalizado for marca in PALABRAS_PROCESAL_EXCLUIR):
        return False, "el nombre indica que es un documento procesal (demanda/memorial/solicitud/etc), no informacion no procesal"

    return None


def es_informacion_no_procesal(nombre_normalizado, contenido_normalizado):
    """
    Rigurosa a propósito: el documento debe SER una tutela, un derecho
    de petición, o un pago oficioso -- no basta con que lo MENCIONE de
    pasada dentro de un documento largo (ej. una demanda que narra en
    su historial que "el demandado interpuso una tutela"). En orden:

    1. El NOMBRE del archivo/asunto ya lo dice (ej. "DERECHO DE
       PETICION ADRESS.pdf") -- señal confiable, se acepta directo.
    2. El nombre trae una marca clara de ser un documento procesal
       (demanda, memorial, solicitud al juzgado, mandamiento, etc, ver
       PALABRAS_PROCESAL_EXCLUIR) -- se descarta, aunque el contenido
       mencione de pasada una tutela/petición/pago oficioso.
    3. Sin marca clara en el nombre: se acepta solo si la frase
       aparece cerca del INICIO del contenido (los primeros ~400
       caracteres -- normalmente el propio encabezado/título del
       documento), no en cualquier parte de un documento de varias
       páginas.
    """
    decision = _decision_solo_por_nombre(nombre_normalizado)
    if decision is not None:
        return decision

    palabras_objetivo = PALABRAS_TIPO_INFORMACION_FUERTES + PALABRAS_PAGO_OFICIOSO
    if any(frase in contenido_normalizado[:400] for frase in palabras_objetivo):
        return True, "tutela/derecho de peticion/pago oficioso (encabezado del contenido)"

    return False, "no parece una tutela, un derecho de peticion, ni un pago oficioso"


def es_auto_terminador(contenido_normalizado):
    """
    True si el contenido da a entender que el proceso judicial NO sigue
    su curso (terminación por pago, retiro/desistimiento de la demanda,
    archivo del proceso, etc) -- no exige que aparezca la palabra
    "AUTO", con que aparezca cualquiera de PALABRAS_PROCESO_NO_CONTINUA
    alcanza.
    """
    return any(frase in contenido_normalizado for frase in PALABRAS_PROCESO_NO_CONTINUA)


# ==================== Gmail (busqueda propia, no la de buscar_faltantes_en_drive.py) ====================


def _texto_plano_html(texto: str) -> str:
    """Quita etiquetas HTML de forma simple, solo para poder buscar palabras clave adentro."""
    return re.sub(r"<[^>]+>", " ", texto or "")


def _leer_correo(mail, id_correo):
    """
    Descarga y parsea UN correo (por su id de IMAP) sobre una conexion
    'mail' YA CONECTADA y con un buzon ya seleccionado -- separado de
    buscar_correo_informacion_no_procesal para que otros scripts (ej.
    clasificar_por_demandado.py, que hace muchas busquedas sobre la
    MISMA conexion en vez de abrir una por termino) puedan reutilizar
    el mismo parseo de MIME sin duplicarlo. Devuelve {"asunto",
    "cuerpo", "adjuntos": [(nombre, bytes), ...], "fecha"}, o None si
    no se pudo descargar.
    """
    typ, msg_datos = mail.fetch(id_correo, "(RFC822)")
    if typ != "OK" or not msg_datos or not msg_datos[0]:
        return None
    mensaje = email.message_from_bytes(msg_datos[0][1])
    asunto = buscador._decodificar_asunto(mensaje.get("Subject", ""))
    fecha_correo = buscador._fecha_del_correo(mensaje)
    cuerpo = ""
    adjuntos = []
    for parte in mensaje.walk():
        nombre_adjunto = parte.get_filename()
        if nombre_adjunto:
            contenido = parte.get_payload(decode=True)
            if contenido:
                adjuntos.append((organizador.sanear_nombre(nombre_adjunto), contenido))
        elif parte.get_content_type() in ("text/html", "text/plain"):
            texto_bruto = parte.get_payload(decode=True)
            if texto_bruto:
                texto = texto_bruto.decode(parte.get_content_charset() or "utf-8", errors="ignore")
                cuerpo += _texto_plano_html(texto) + " "
    return {"asunto": asunto, "cuerpo": cuerpo, "adjuntos": adjuntos, "fecha": fecha_correo}


def buscar_correo_informacion_no_procesal(usuario: str, app_password: str, termino: str):
    """
    Busca en TODO el Gmail (via X-GM-RAW, igual que buscar_en_correo de
    buscar_faltantes_en_drive.py -- es la misma busqueda que escribir
    'termino' en la barra de busqueda de Gmail, sobre "Todos los
    mensajes") los correos que mencionen 'termino'. A diferencia de esa
    función (que solo mira enlaces de Drive y adjuntos .zip), esta
    devuelve TODO lo que hace falta para clasificar un correo como
    tutela/derecho de petición/pago oficioso: asunto, cuerpo de texto
    (HTML ya limpiado), y CUALQUIER adjunto (no solo .zip). Devuelve
    [{"asunto", "cuerpo", "adjuntos": [(nombre, bytes), ...], "fecha"}, ...].
    """
    resultados = []
    with imaplib.IMAP4_SSL("imap.gmail.com") as mail:
        mail.login(usuario, app_password)
        mail.select('"[Gmail]/All Mail"', readonly=True)
        typ, datos = mail.search(None, "X-GM-RAW", f'"{termino}"')
        if typ != "OK" or not datos or not datos[0]:
            return resultados
        for id_correo in datos[0].split():
            correo = _leer_correo(mail, id_correo)
            if correo is not None:
                resultados.append(correo)
    return resultados


def _extraer_pdfs_de_zip(contenido: bytes, destino: Path) -> int:
    extraidos = 0
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as archivo_zip:
            for info in archivo_zip.infolist():
                if info.is_dir() or Path(info.filename).suffix.lower() not in (".pdf", ".docx"):
                    continue
                nombre_seguro = organizador.sanear_nombre(Path(info.filename).name)
                ruta = buscador._ruta_archivo_libre(destino, nombre_seguro)
                with archivo_zip.open(info) as origen, open(organizador._ruta_larga_segura(str(ruta)), "wb") as f:
                    f.write(origen.read())
                extraidos += 1
    except zipfile.BadZipFile:
        pass
    return extraidos


def _guardar_correo(correo: dict, destino: Path) -> int:
    """
    Guarda los adjuntos PDF/DOCX del correo (extrayendo los de dentro
    de un .zip si aplica). Si no trae ningun adjunto util, guarda el
    asunto + cuerpo como un .txt simple, para no perder la informacion
    (ej. un correo de pago oficioso en texto plano, sin adjuntos).
    """
    _crear_carpeta(destino)
    guardados = 0
    for nombre_adjunto, contenido in correo["adjuntos"]:
        extension = Path(nombre_adjunto).suffix.lower()
        if extension == ".zip":
            guardados += _extraer_pdfs_de_zip(contenido, destino)
        elif extension in (".pdf", ".docx"):
            ruta = buscador._ruta_archivo_libre(destino, nombre_adjunto)
            with open(organizador._ruta_larga_segura(str(ruta)), "wb") as f:
                f.write(contenido)
            guardados += 1

    if guardados == 0:
        fecha_texto = correo["fecha"].isoformat() if correo["fecha"] else "sin_fecha"
        nombre_txt = organizador.sanear_nombre(f"{fecha_texto} - {correo['asunto'] or 'correo'}.txt")
        ruta = buscador._ruta_archivo_libre(destino, nombre_txt)
        with open(organizador._ruta_larga_segura(str(ruta)), "w", encoding="utf-8") as f:
            f.write(f"Asunto: {correo['asunto']}\nFecha: {correo['fecha']}\n\n{correo['cuerpo']}")
        guardados = 1
    return guardados


def buscar_correo_global_informacion_no_procesal(usuario: str, app_password: str):
    """
    Busca en TODO Gmail, UNA SOLA VEZ para toda la corrida (no una vez
    por proceso), los correos que mencionen alguna de las frases de
    PALABRAS_TIPO_INFORMACION_FUERTES o PALABRAS_PAGO_OFICIOSO --
    exactamente las que usa es_informacion_no_procesal. Se busca por
    TIPO de correo (tutela/derecho de peticion/pago oficioso), no por radicado:
    un correo de este tipo no siempre menciona el radicado exacto tal
    como Gmail lo indexaria, así que primero se encuentra por su
    contenido y DESPUES se empareja con el proceso correcto por
    radicado, cuenta, o nombre del demandado -- ver
    _procesos_que_coinciden_con_correo. Devuelve la lista de correos
    (sin duplicados, [{"asunto", "cuerpo", "adjuntos", "fecha"}, ...]).
    """
    vistos = {}
    for frase in PALABRAS_TIPO_INFORMACION_FUERTES + PALABRAS_PAGO_OFICIOSO:
        try:
            correos = buscar_correo_informacion_no_procesal(usuario, app_password, frase)
        except Exception as error:
            logging.error("[Correo] Fallo buscando '%s' en Gmail: %s", frase, error)
            continue
        for correo in correos:
            vistos[(correo["asunto"], correo["fecha"])] = correo
    return list(vistos.values())



# Un demandado cuyo nombre se reduce a UNA sola palabra significativa
# (tras quitar PALABRAS_GENERICAS_DEMANDADO) es una señal demasiado
# debil para cruzar solo -- puede ser un dato mal diligenciado en el
# Excel (ej. una palabra suelta en vez de un nombre completo), y
# cualquier archivo/correo que por casualidad use esa misma palabra
# (aunque sea un termino administrativo comun) terminaria
# "coincidiendo" con ese proceso. Por eso el cruce por demandado exige
# como minimo este numero de palabras significativas.
MIN_PALABRAS_DEMANDADO_PARA_CRUZAR = 2

# El nombre del demandado solo se busca en los primeros N caracteres
# del contenido (encabezado/caratula del documento -- en Colombia
# practicamente todo documento judicial identifica las partes ahi
# mismo, al principio) -- NO en el documento completo. Un PDF largo
# (o uno que junta varios casos en un solo archivo escaneado) puede
# mencionar OTROS nombres mas adelante (ej. el nombre de un municipio
# que tambien es el del JUZGADO que emite el documento, o un caso
# distinto en las paginas siguientes) que no tienen nada que ver con
# el demandado real de ESE documento. El radicado y la cuenta SI se
# buscan en el documento completo (son coincidencias mucho mas
# precisas por si solas, no hace falta restringirlas).
VENTANA_DEMANDADO_CARACTERES = 3000

# El MEMBRETE del juzgado (su propio nombre y direccion -- ej. "JUZGADO
# PROMISCUO MUNICIPAL DE CANTAGALLO, BOLIVAR", seguido de la
# direccion/telefono/correo del despacho) es literalmente lo PRIMERO
# que traen estos documentos, antes incluso de identificar a las
# partes -- y casi siempre incluye el nombre del MUNICIPIO donde queda
# el juzgado, que puede coincidir por pura casualidad con el demandado
# de OTRO proceso (ej. un auto del "Juzgado de Cantagallo" que no es
# contra el Municipio de Cantagallo, sino contra otra persona, pero
# menciona "Cantagallo" en su propio membrete). Por eso estas lineas se
# quitan ANTES de buscar el demandado -- nunca deben contar como si
# fueran parte del proceso.
_PATRON_MEMBRETE_JUZGADO = re.compile(
    r"(?:^|\n)[^\n]*(?:"
    r"JUZGADO|RAMA JUDICIAL|CONSEJO SECCIONAL DE LA JUDICATURA|"
    r"@CENDOJ|RAMAJUDICIAL\.GOV\.CO|"
    r"\b(?:CARRERA|CALLE|AVENIDA|TRANSVERSAL|DIAGONAL|BARRIO)\b[^\n]*\d|"
    r"\bTELEFONO\b|\bCORREO\s*:"
    r")[^\n]*",
    re.IGNORECASE,
)


_MESES_ESPANOL = (
    "ENERO|FEBRERO|MARZO|ABRIL|MAYO|JUNIO|JULIO|AGOSTO|SEPTIEMBRE|OCTUBRE|NOVIEMBRE|DICIEMBRE"
)

# La FECHA DE CIERRE/FIRMA de un auto colombiano casi siempre se
# escribe "<MUNICIPIO donde queda el juzgado>, <dia> DE <mes> DE
# <anio>" (ej. "SABANA DE TORRES, DIEZ (10) DE JULIO DE 2025" o
# "PUERTO WILCHES (S), 20 DE JUNIO DE 2023") -- es la MISMA trampa que
# el membrete (el nombre del municipio, no del demandado), pero
# aparece al FINAL del documento en vez de al principio. Se reconoce
# por el patron fijo "DE <mes> DE" (los 12 meses son una lista
# cerrada) y se quita el nombre del lugar justo antes de la coma.
_PATRON_FECHA_CIERRE = re.compile(
    rf"[A-ZÑ][A-ZÑ ]{{1,35}}(?:\s*\([A-ZÑ]\))?,\s*"
    rf"(?:[A-ZÑ()0-9\s]{{0,25}}\bDE\s+)?"
    rf"(?:{_MESES_ESPANOL})\s+DE\s+"
    rf"[A-ZÑ()0-9\s]{{0,25}}"
)


# El PDF extrae el texto respetando el SALTO DE LINEA VISUAL del
# documento (el "word wrap" de cada renglon impreso), no las oraciones
# reales -- una frase del membrete/firma puede quedar partida justo a
# la mitad (ej. "...Municipio de Puerto\nWilches (S). ,"), y como
# _PATRON_MEMBRETE_JUZGADO/_PATRON_FECHA_CIERRE trabajan LINEA POR
# LINEA (delimitadas por "\n"), la mitad que cae en el renglon
# siguiente se escapaba del filtro -- "WILCHES" sobrevivia solo,
# separado de "PUERTO", y terminaba pareciendo un demandado real. Por
# eso ANTES de aplicar esos dos filtros se unen los saltos de linea
# "sueltos" (los que van DENTRO de un mismo parrafo) en un espacio.
# Se deja INTACTO el salto de linea (no se une) en dos casos, para no
# pegar por error dos cosas que NO son el mismo parrafo:
#   - Una linea en blanco (separador real de parrafo).
#   - Una linea siguiente que empieza como una ETIQUETA de campo (ej.
#     "RADICADO:", "DEMANDANTE:", "DEMANDADO:", "CONSTANCIA:") -- estos
#     campos suelen venir uno debajo del otro sin linea en blanco entre
#     ellos, y unirlos con el parrafo de arriba (que a veces SI
#     contiene "JUZGADO" en su propio texto) terminaba borrando por
#     error el campo DEMANDADO/RADICADO real del documento.
#
# Esta union SOLO se aplica antes de _PATRON_MEMBRETE_JUZGADO (que es
# el que de verdad la necesita) -- _PATRON_FECHA_CIERRE se aplica ANTES
# de unir nada, sobre el texto tal cual lo separa el PDF en renglones:
# si se uniera primero, la fecha de cierre podia "alcanzar hacia atras"
# por encima de un salto de linea real y borrar de paso el nombre del
# demandado que estuviera juntico arriba (ej. "DEMANDADO: MUNICIPIO...\n
# CANTAGALLO, DIEZ DE JULIO DE 2025." -- unido en una sola linea, el
# patron de fecha alcanzaba a comerse tambien el nombre del demandado).
_PATRON_SALTO_DE_LINEA_SUELTO = re.compile(
    r"\n(?!\s*\n)(?!\s*[A-ZÑ][A-ZÑ .]{1,25}:)"
)


def _unir_lineas_de_un_mismo_parrafo(texto: str) -> str:
    return _PATRON_SALTO_DE_LINEA_SUELTO.sub(" ", texto)


def _quitar_membrete_juzgado(texto_normalizado: str) -> str:
    """Quita la fecha de cierre/firma, y el membrete/direccion/contacto del juzgado (ambos suelen traer el nombre del MUNICIPIO, no del demandado) -- ver _PATRON_FECHA_CIERRE y _PATRON_MEMBRETE_JUZGADO."""
    sin_fecha_cierre = _PATRON_FECHA_CIERRE.sub(" ", texto_normalizado)
    texto_unido = _unir_lineas_de_un_mismo_parrafo(sin_fecha_cierre)
    return _PATRON_MEMBRETE_JUZGADO.sub(" ", texto_unido)


def _palabra_demandado_coincide(texto_normalizado, palabra):
    """
    True si 'palabra' aparece en 'texto_normalizado' como PALABRA
    COMPLETA (no pegada a otras letras) -- a diferencia de
    buscador._nombre_coincide (que solo protege límites de DIGITOS, no
    de letras, porque para radicados/cuentas eso es lo que hace falta),
    esto evita que una palabra del demandado "encaje" por casualidad
    DENTRO de un nombre de archivo sin espacios (ej. "AUTO" adentro de
    "11AUTOTERMINAPROCESO.pdf", que no tiene relación con ningún
    demandado que se llame o incluya "AUTO").
    """
    patron = re.compile(rf"(?<![A-ZÑ]){re.escape(palabra)}(?![A-ZÑ])")
    return patron.search(texto_normalizado) is not None


def _indexar_procesos_para_correo(procesos):
    """
    Indices para emparejar un correo/archivo con el/los proceso(s) al
    que corresponde, por radicado, por cuenta, o por el NOMBRE COMPLETO
    (todas sus palabras significativas) del demandado. Un mismo correo
    puede corresponder a mas de un proceso (ej. un proceso "acumulado"
    con varias cuentas bajo un radicado, o el mismo demandado con mas
    de un proceso activo).
    """
    por_radicado, por_cuenta, por_demandado = {}, {}, {}
    for proceso in procesos:
        if proceso["radicado"]:
            por_radicado.setdefault(proceso["radicado"], []).append(proceso)
        for cuenta in proceso["cuentas"]:
            if buscador._cuenta_es_valida_para_buscar(cuenta):
                por_cuenta.setdefault(cuenta, []).append(proceso)
        for demandado in proceso["demandados"]:
            palabras = frozenset(buscador._palabras_significativas(demandado))
            if len(palabras) >= MIN_PALABRAS_DEMANDADO_PARA_CRUZAR:
                por_demandado.setdefault(palabras, []).append(proceso)
    return por_radicado, por_cuenta, por_demandado


def _procesos_que_coinciden_con_correo(contenido_normalizado, indices):
    """
    Devuelve los procesos (sin duplicados) a los que este correo/archivo
    corresponde: basta con que coincida el radicado, la cuenta, O el
    demandado (cualquiera de los tres, no hace falta que coincidan
    todos).

    Para el RADICADO se reconoce tanto el número PLANO (23 dígitos
    seguidos) como el escrito CON separadores (guiones, puntos, o
    espacios entre sus grupos -- ej. "68001-40-03-001-2024-00050-00",
    el formato más común dentro del texto de un documento real -- ver
    cruce_excel._radicados_en_texto).

    Para el DEMANDADO se exige que aparezcan TODAS sus palabras
    significativas, COMO PALABRA COMPLETA, dentro de los primeros
    VENTANA_DEMANDADO_CARACTERES del texto UNA VEZ QUITADO EL MEMBRETE
    DEL JUZGADO (ver _quitar_membrete_juzgado) -- no en el documento
    completo, y sin contar el nombre/dirección/correo del despacho que
    emite el documento. Sin esto, el nombre del MUNICIPIO donde queda
    el juzgado (que casi siempre aparece en su propio membrete, ej.
    "JUZGADO PROMISCUO MUNICIPAL DE CANTAGALLO, BOLIVAR") se confundía
    con el demandado de OTRO proceso que sí tenga ese municipio como
    demandado real, aunque el documento no tuviera nada que ver con él
    -- generaba coincidencias masivas y completamente falsas.
    """
    por_radicado, por_cuenta, por_demandado = indices
    encontrados = {}

    # El radicado dentro de un documento casi nunca aparece "plano" --
    # normalmente trae guiones, puntos, o espacios entre sus grupos de
    # digitos (ej. "68001-40-03-001-2024-00050-00"). cruce_excel._radicados_en_texto
    # reconoce esa forma (y la plana) y devuelve el numero ya limpio de
    # separadores, para poder comparar directo contra el radicado del
    # Excel -- buscar la cadena plana tal cual, sin esto, se perdia
    # cualquier radicado que el documento no repitiera exactamente sin
    # separadores.
    radicados_del_texto = set(cruce_excel._radicados_en_texto(contenido_normalizado))
    for radicado, procesos in por_radicado.items():
        coincide_radicado = radicado in radicados_del_texto or any(
            buscador._nombre_coincide(contenido_normalizado, t) for t in buscador.radicados_cortos(radicado)
        )
        if coincide_radicado:
            for proceso in procesos:
                encontrados[proceso["nombre_carpeta"]] = proceso

    for cuenta, procesos in por_cuenta.items():
        if buscador._nombre_coincide(contenido_normalizado, cuenta):
            for proceso in procesos:
                encontrados[proceso["nombre_carpeta"]] = proceso

    encabezado = _quitar_membrete_juzgado(contenido_normalizado)[:VENTANA_DEMANDADO_CARACTERES]
    for palabras, procesos in por_demandado.items():
        if all(_palabra_demandado_coincide(encabezado, palabra) for palabra in palabras):
            for proceso in procesos:
                encontrados[proceso["nombre_carpeta"]] = proceso

    return list(encontrados.values())


def procesar_correos_no_procesal(correos, procesos_con_radicado, carpetas_existentes):
    """
    Clasifica los correos ya encontrados globalmente (ver
    buscar_correo_global_informacion_no_procesal) y los adjunta a la
    carpeta de CADA proceso con el que coincidan por radicado, cuenta o
    demandado (ver _procesos_que_coinciden_con_correo). Siempre exige
    ademas que el correo mencione a ESSA/Electrificadora de Santander.
    """
    indices = _indexar_procesos_para_correo(procesos_con_radicado)
    sin_proceso = 0
    adjuntados = 0

    for correo in correos:
        asunto_norm = buscador._normalizar_para_comparar(correo["asunto"])
        contenido_norm = buscador._normalizar_para_comparar(
            correo["asunto"] + " " + correo["cuerpo"] + " " + " ".join(n for n, _ in correo["adjuntos"])
        )
        es_no_procesal, motivo = es_informacion_no_procesal(asunto_norm, contenido_norm)
        if not es_no_procesal:
            continue

        procesos_coincidentes = _procesos_que_coinciden_con_correo(contenido_norm, indices)
        if not procesos_coincidentes:
            sin_proceso += 1
            logging.info(
                "   [Correo] '%s' parece %s, pero no coincide con el radicado/cuenta/demandado de ningun "
                "proceso -- no se pudo emparejar, se omite.", correo["asunto"], motivo,
            )
            continue

        tiene_essa = any(buscador._nombre_coincide(contenido_norm, t) for t in buscador.TERMINOS_DEMANDANTE_VALIDO)
        if not tiene_essa:
            logging.info(
                "   [Correo] se omite '%s': no se confirmo que sea de ESSA/Electrificadora de Santander.",
                correo["asunto"],
            )
            continue

        for proceso in procesos_coincidentes:
            numero, nombre_carpeta = proceso["numero"], proceso["nombre_carpeta"]
            destino = Path(CARPETA_PROCESOS) / nombre_carpeta

            if MODO_PRUEBA:
                logging.info(
                    "[SIMULACION] Proceso %s: subiria el correo '%s' (%s) a '%s'.",
                    numero, correo["asunto"], motivo, nombre_carpeta,
                )
                adjuntados += 1
                continue

            if nombre_carpeta not in carpetas_existentes:
                _crear_carpeta(destino)
                carpetas_existentes.add(nombre_carpeta)
                logging.info("[Creada] Proceso %s: carpeta '%s' (encontrada por correo).", numero, nombre_carpeta)

            guardados = _guardar_correo(correo, destino)
            logging.info(
                "[Descargado] Proceso %s: correo '%s' (%s) -> %d archivo(s) en %s",
                numero, correo["asunto"], motivo, guardados, nombre_carpeta,
            )
            adjuntados += guardados

    logging.info(
        "[Correo] %d correo(s) de informacion no procesal adjuntados a algun proceso; %d no se pudieron "
        "emparejar con ningun proceso conocido.", adjuntados, sin_proceso,
    )


# ==================== Carpetas en disco ====================


def _crear_carpeta(destino: Path):
    """
    mkdir a prueba de rutas largas de Windows -- sin el prefijo especial
    (ver organizador._ruta_larga_segura), una carpeta cuyo camino
    completo (CARPETA_PROCESOS + nombre) supere ~260 caracteres falla
    con un OSError silencioso que hacia que TODA la fila se saltara
    (quedaba solo un "[Error] ... fallo y se omite" generico en el log,
    sin crear la carpeta ni buscar nada para ese proceso).
    """
    os.makedirs(organizador._ruta_larga_segura(str(destino)), exist_ok=True)


def crear_todas_las_carpetas(procesos, carpetas_existentes, etiqueta):
    """
    Crea de una vez TODAS las carpetas que hagan falta para 'procesos'
    (activos/suspendidos/etc, o terminados) -- se corre ANTES de
    empezar a buscar documentos en Drive/Gmail para que las 1223
    carpetas queden visibles en el disco desde el principio, en vez de
    ir apareciendo intercaladas a medida que el script busca contenido
    (que puede tardar horas para los ~925 procesos con radicado).

    Los procesos "TERMINADO POR AUTO" no reciben carpeta propia: todos
    comparten CARPETA_TERMINADOS_POR_AUTO, que se crea UNA sola vez (su
    auto se guarda ahi renombrado, ver procesar_terminado).
    """
    creadas = 0
    for proceso in procesos:
        nombre_carpeta = proceso["carpeta_destino"]
        if nombre_carpeta in carpetas_existentes:
            continue
        destino = Path(CARPETA_PROCESOS) / nombre_carpeta
        filas = ", ".join(str(f) for f in proceso["filas_excel"])
        if MODO_PRUEBA:
            logging.info(
                "[SIMULACION] Proceso %s (%s, fila(s) %s): se crearia la carpeta '%s'.",
                proceso["numero"], proceso["estado"], filas, nombre_carpeta,
            )
            # En MODO_PRUEBA no se crea nada, asi que se anota igual para
            # no repetir la misma linea por cada uno de los cientos de
            # procesos que comparten CARPETA_TERMINADOS_POR_AUTO.
            carpetas_existentes.add(nombre_carpeta)
        else:
            _crear_carpeta(destino)
            carpetas_existentes.add(nombre_carpeta)
            logging.info(
                "[Creada] Proceso %s (%s, fila(s) %s): carpeta '%s'.",
                proceso["numero"], proceso["estado"], filas, nombre_carpeta,
            )
        creadas += 1
    logging.info("%s: %d carpeta(s) %s.", etiqueta, creadas, "simuladas (MODO_PRUEBA activo)" if MODO_PRUEBA else "creadas")


def _contenido_normalizado_local(ruta: Path) -> str:
    """
    Nombre + texto de un archivo que YA esta en el disco, normalizado
    igual que _info_documento (que lee el de Drive) -- para poder
    aplicarle es_auto_terminador a lo que dejaron corridas anteriores.
    """
    extension = ruta.suffix.lower()
    ruta_segura = organizador._ruta_larga_segura(str(ruta))
    texto = ""
    if extension == ".pdf":
        texto = organizador.texto_de_pdf(ruta_segura)
    elif extension == ".docx" and organizador.docx is not None:
        texto = organizador.texto_de_docx(ruta_segura)
    return buscador._normalizar_para_comparar(ruta.name + " " + texto)


def mover_autos_de_carpetas_antiguas(terminados, carpetas_existentes):
    """
    Corridas ANTERIORES (de antes de CARPETA_TERMINADOS_POR_AUTO) dejaron
    cada auto dentro de la carpeta propia del proceso ("245. TERMINADO
    POR AUTO"), con el nombre que traia en Drive. Este paso los recoge y
    los deja donde van ahora: dentro de CARPETA_TERMINADOS_POR_AUTO y
    renombrados con el numero del proceso ("245. TERMINADO POR AUTO.pdf").

    Solo toca las carpetas de procesos "TERMINADO POR AUTO", y dentro de
    ellas SOLO los documentos que de verdad parecen el auto que termina
    el proceso (mismo criterio que la descarga: es_auto_terminador sobre
    el nombre y el texto del archivo). Cualquier otra cosa que alguien
    haya dejado ahi (una petición, una respuesta, un anexo suelto) se
    queda donde esta: renombrarla "245. TERMINADO POR AUTO.pdf" seria
    mentir sobre lo que es.

    Nunca pisa un archivo que ya exista en el destino (usa "..._2",
    "..._3", igual que una descarga nueva) y nunca borra un documento:
    la carpeta vieja se borra UNICAMENTE si quedo completamente vacia
    despues de mover el auto. Respeta MODO_PRUEBA.
    """
    carpeta_comun = Path(CARPETA_PROCESOS) / CARPETA_TERMINADOS_POR_AUTO
    movidos = 0
    carpetas_vaciadas = 0

    for proceso in terminados:
        if not proceso["nombre_archivo"]:
            continue  # terminado por pago/contrato/no inicio: conserva su carpeta propia
        nombre_viejo = proceso["nombre_carpeta"]
        if nombre_viejo not in carpetas_existentes:
            continue
        carpeta_vieja = Path(CARPETA_PROCESOS) / nombre_viejo

        try:
            archivos = [ruta for ruta in carpeta_vieja.iterdir() if ruta.is_file()]
        except OSError as error:
            logging.warning("Proceso %s: no se pudo leer la carpeta '%s', se omite: %s", proceso["numero"], nombre_viejo, error)
            continue

        for ruta in archivos:
            if not es_auto_terminador(_contenido_normalizado_local(ruta)):
                logging.info(
                    "   (se deja '%s' en '%s': no parece el auto que termina el proceso %s)",
                    ruta.name, nombre_viejo, proceso["numero"],
                )
                continue

            extension = ruta.suffix or ".pdf"
            nombre_nuevo = organizador.sanear_nombre(f"{proceso['nombre_archivo']}{extension}")
            if MODO_PRUEBA:
                logging.info(
                    "[SIMULACION] Proceso %s: movería '%s' de '%s' a '%s' como '%s'.",
                    proceso["numero"], ruta.name, nombre_viejo, CARPETA_TERMINADOS_POR_AUTO, nombre_nuevo,
                )
                movidos += 1
                continue
            _crear_carpeta(carpeta_comun)
            carpetas_existentes.add(CARPETA_TERMINADOS_POR_AUTO)
            ruta_final = buscador._ruta_archivo_libre(carpeta_comun, nombre_nuevo)
            try:
                os.replace(
                    organizador._ruta_larga_segura(str(ruta)),
                    organizador._ruta_larga_segura(str(ruta_final)),
                )
            except OSError as error:
                logging.warning(
                    "Proceso %s: no se pudo mover '%s' (¿archivo abierto, o bloqueado por OneDrive/antivirus?); "
                    "se deja donde esta y se sigue: %s", proceso["numero"], ruta.name, error,
                )
                continue
            logging.info("[Movido] Proceso %s: '%s' -> %s", proceso["numero"], ruta.name, ruta_final)
            movidos += 1

        if MODO_PRUEBA:
            continue
        resto = list(carpeta_vieja.iterdir())
        if resto:
            # Lo normal aqui es la subcarpeta CORREOS (correos que le
            # siguieron llegando al proceso ya terminado, ver
            # carpeta_de_correos) -- eso no es nada que revisar. Solo se
            # advierte si quedaron archivos sueltos que no se pudieron mover.
            if all(ruta.is_dir() for ruta in resto):
                logging.info(
                    "Proceso %s: se conserva la carpeta '%s' porque tiene subcarpeta(s) (%s).",
                    proceso["numero"], nombre_viejo, ", ".join(sorted(r.name for r in resto)),
                )
            else:
                logging.info(
                    "Proceso %s: se conserva la carpeta '%s' con %d archivo(s) que no son el auto "
                    "(o que no se pudieron mover) -- revisala cuando puedas.",
                    proceso["numero"], nombre_viejo, sum(1 for r in resto if r.is_file()),
                )
            continue
        try:
            os.rmdir(organizador._ruta_larga_segura(str(carpeta_vieja)))
        except OSError as error:
            logging.warning("Proceso %s: no se pudo borrar la carpeta vacia '%s': %s", proceso["numero"], nombre_viejo, error)
            continue
        carpetas_existentes.discard(nombre_viejo)
        carpetas_vaciadas += 1

    if movidos or carpetas_vaciadas:
        logging.info(
            "[Terminados por auto] %d auto(s) %s a '%s'; %d carpeta(s) vieja(s) vacia(s) borrada(s).",
            movidos, "se moverian (MODO_PRUEBA activo)" if MODO_PRUEBA else "movidos",
            CARPETA_TERMINADOS_POR_AUTO, carpetas_vaciadas,
        )


def _listar_carpetas_existentes():
    carpeta_raiz = Path(CARPETA_PROCESOS)
    _crear_carpeta(carpeta_raiz)
    return {
        d.name for d in carpeta_raiz.iterdir()
        if d.is_dir() and d.name not in cruce_excel.CARPETAS_A_IGNORAR
    }


def _cargar_procesados_anteriormente() -> set:
    """Carpetas ya revisadas por completo en una corrida anterior (ver ARCHIVO_PROCESADOS) -- se omiten esta vez."""
    if not os.path.exists(ARCHIVO_PROCESADOS):
        return set()
    with open(ARCHIVO_PROCESADOS, encoding="utf-8") as f:
        return {linea.strip() for linea in f if linea.strip()}


_lock_archivo_procesados = threading.Lock()


def _marcar_como_procesado(nombre_carpeta: str):
    """
    Registra 'nombre_carpeta' como ya revisada, para que la proxima
    corrida la omita. No hace nada en MODO_PRUEBA. Con hilos en paralelo
    (ver NUM_HILOS) varios procesos pueden terminar casi al mismo
    tiempo, asi que el escribir al archivo va protegido con un lock para
    que dos hilos no se pisen escribiendo a la vez.
    """
    if MODO_PRUEBA:
        return
    with _lock_archivo_procesados:
        with open(ARCHIVO_PROCESADOS, "a", encoding="utf-8") as f:
            f.write(nombre_carpeta + "\n")


def _nombre_destino(archivo, nombre_base: str) -> str:
    """
    Nombre final del archivo descargado cuando se va a RENOMBRAR (ej.
    "245. TERMINADO POR AUTO.pdf"): 'nombre_base' + la extension que le
    corresponde al documento. Los archivos nativos de Google (Doc,
    Sheet, Slide) se exportan a PDF, asi que su extension siempre es
    ".pdf"; para el resto se conserva la del original (aqui casi siempre
    .pdf tambien -- ver buscador._es_pdf_o_exportable). La extension es
    obligatoria: sin ella, un nombre como "245. TERMINADO POR AUTO"
    haria que Path lo interpretara como extension ". TERMINADO POR AUTO".
    """
    if archivo["mimeType"] in buscador.MIME_EXPORTAR:
        extension = buscador.MIME_EXPORTAR[archivo["mimeType"]][0]
    else:
        extension = Path(archivo.get("name", "")).suffix or ".pdf"
    return organizador.sanear_nombre(f"{nombre_base}{extension}")


def _descargar_archivo(servicio, archivo, destino: Path, nombre_base: str = None) -> Path:
    """
    Descarga 'archivo' de Drive dentro de 'destino'. Con 'nombre_base'
    se guarda RENOMBRADO (ver _nombre_destino); sin el, conserva su
    nombre original de Drive. Nunca pisa un archivo que ya exista: si el
    nombre esta ocupado (ej. un proceso con auto de primera y de segunda
    instancia) se guarda como "..._2", "..._3", etc.
    """
    _crear_carpeta(destino)
    nombre_seguro = _nombre_destino(archivo, nombre_base) if nombre_base else organizador.sanear_nombre(archivo["name"])
    ruta_local = buscador._ruta_archivo_libre(destino, nombre_seguro)
    if archivo["mimeType"] in buscador.MIME_EXPORTAR:
        buscador._exportar_google_doc(servicio, archivo["id"], archivo["mimeType"], ruta_local)
    else:
        buscador._descargar_archivo_binario(servicio, archivo["id"], ruta_local)
    return ruta_local


# ==================== Procesamiento por fila ====================


def procesar_con_radicado(proceso):
    """
    Filas que NO son terminadas (activo, suspendido, reorganizacion,
    remitida a castigo/prepago, etc). La carpeta ya se creo antes (ver
    crear_todas_las_carpetas). La busqueda en Gmail para estas filas
    tampoco se hace aqui -- se hace UNA vez para todas al final, ver
    procesar_correos_no_procesal. Corre en un hilo propio (ver
    _servicio_del_hilo) para que varios procesos se busquen en Drive al
    mismo tiempo -- ver NUM_HILOS.
    """
    servicio = _servicio_del_hilo()
    numero, radicado = proceso["numero"], proceso["radicado"]
    nombre_carpeta = proceso["carpeta_destino"]
    destino = Path(CARPETA_PROCESOS) / nombre_carpeta

    archivos = _buscar_archivos(servicio, radicado, proceso["cuentas"])

    subidos = 0
    for archivo, carpeta in archivos:
        nombre_norm = buscador._normalizar_para_comparar(archivo.get("name", ""))
        decision_por_nombre = _decision_solo_por_nombre(nombre_norm)
        if decision_por_nombre is not None and decision_por_nombre[0] is False:
            # El nombre ya lo descarta (demanda/memorial/etc) -- ni
            # siquiera hace falta descargarlo para saberlo.
            logging.info("   (se omite '%s' del proceso %s: %s)", archivo["name"], numero, decision_por_nombre[1])
            continue

        nombre_norm, contenido_norm = _info_documento(servicio, archivo)
        seguro, motivo = _archivo_es_seguro(carpeta.get("name", ""), archivo, contenido_norm, proceso["demandados"])
        if not seguro:
            logging.info("   (se omite '%s' del proceso %s: %s)", archivo["name"], numero, motivo)
            continue

        es_no_procesal, motivo = decision_por_nombre if decision_por_nombre is not None else es_informacion_no_procesal(nombre_norm, contenido_norm)
        if not es_no_procesal:
            logging.info("   (se omite '%s' del proceso %s: %s)", archivo["name"], numero, motivo)
            continue

        if MODO_PRUEBA:
            logging.info("[SIMULACION] Proceso %s: subiria '%s' (%s) a '%s'.", numero, archivo["name"], motivo, nombre_carpeta)
        else:
            ruta = _descargar_archivo(servicio, archivo, destino)
            logging.info("[Descargado] Proceso %s: '%s' -> %s", numero, archivo["name"], ruta)
        subidos += 1

    if subidos == 0:
        logging.info(
            "Proceso %s (%s, radicado %s): no se encontro informacion no procesal en Drive todavia "
            "(el correo se revisa aparte, al final de la corrida).", numero, proceso["estado"], radicado,
        )


def procesar_terminado(proceso, pendientes) -> bool:
    """
    La carpeta ya se creo antes (ver crear_todas_las_carpetas). Devuelve
    True si se encontro el documento que termina el proceso (para que
    procesar() lo marque como revisado y no lo vuelva a buscar en la
    proxima corrida) -- si queda pendiente, devuelve False para que se
    reintente la proxima vez. Corre en un hilo propio (ver
    _servicio_del_hilo) -- ver NUM_HILOS.

    Para los "TERMINADO POR AUTO", el auto no se deja con el nombre que
    traia en Drive: se guarda como "<numero del Excel>. <ESTADO
    PROCESAL>.pdf" (ej. "245. TERMINADO POR AUTO.pdf") dentro de la
    carpeta comun CARPETA_TERMINADOS_POR_AUTO. El resto de terminados
    (pago/contrato/no inicio) sigue igual que antes: su propia carpeta y
    el nombre original del documento.
    """
    servicio = _servicio_del_hilo()
    numero, estado, radicado = proceso["numero"], proceso["estado"], proceso["radicado"]
    nombre_carpeta = proceso["carpeta_destino"]
    nombre_archivo = proceso["nombre_archivo"]
    destino = Path(CARPETA_PROCESOS) / nombre_carpeta

    if not _terminos_busqueda(radicado, proceso["cuentas"]):
        pendientes.append(proceso)
        logging.warning("Proceso %s (%s): no hay radicado ni cuenta valida para buscar en Drive, queda pendiente.", numero, estado)
        return False

    archivos = _buscar_archivos(servicio, radicado, proceso["cuentas"])

    encontrado = False
    for archivo, carpeta in archivos:
        _nombre_norm, contenido_norm = _info_documento(servicio, archivo)
        seguro, motivo = _archivo_es_seguro(carpeta.get("name", ""), archivo, contenido_norm, proceso["demandados"])
        if not seguro:
            logging.info("   (se omite '%s' del proceso %s: %s)", archivo["name"], numero, motivo)
            continue

        if not es_auto_terminador(contenido_norm):
            continue

        if MODO_PRUEBA:
            logging.info(
                "[SIMULACION] Proceso %s (%s): subiria '%s' (termina el proceso) a '%s'%s.",
                numero, estado, archivo["name"], nombre_carpeta,
                f", renombrado como '{_nombre_destino(archivo, nombre_archivo)}'" if nombre_archivo else "",
            )
        else:
            ruta = _descargar_archivo(servicio, archivo, destino, nombre_archivo)
            logging.info("[Descargado] Proceso %s (%s): '%s' -> %s", numero, estado, archivo["name"], ruta)
        encontrado = True
        # No se corta el ciclo: puede haber mas de un documento relevante (ej. primera y segunda instancia).

    if not encontrado:
        pendientes.append(proceso)
        logging.warning("Proceso %s (%s): no se encontro el documento que termina el proceso, queda pendiente.", numero, estado)

    return encontrado


# ==================== Orquestacion ====================


def procesar():
    if not RUTA_EXCEL_CONTROL or not os.path.exists(RUTA_EXCEL_CONTROL):
        logging.error("No se encontro el Excel configurado en RUTA_EXCEL_CONTROL: %r", RUTA_EXCEL_CONTROL)
        return

    procesos = leer_procesos_control()
    con_radicado, terminados, sin_estado = clasificar_procesos(procesos)
    logging.info(
        "Excel: %d fila(s) con numero de proceso -- %d activo/suspendido/reorganizacion/remitida/etc "
        "(carpeta con radicado, o con ESTADO PROCESAL si todavia no lo tiene), %d terminadas (pago/auto/"
        "contrato/no inicio), %d sin ESTADO PROCESAL diligenciado.",
        len(procesos), len(con_radicado), len(terminados), len(sin_estado),
    )

    servicio = autenticar_drive_o_none()
    if servicio is None:
        logging.error("No se puede continuar sin conexion a Google Drive.")
        return

    credenciales_correo = None
    if BUSCAR_EN_CORREO:
        credenciales_correo = organizador.leer_credenciales()
        if not credenciales_correo:
            logging.warning(
                "[Correo] No hay %s (o le faltan datos); se omite la busqueda en Gmail.",
                organizador.ARCHIVO_CREDENCIALES,
            )
        else:
            logging.info(
                "[Correo] Credenciales encontradas (%s) -- se buscara en Gmail por tutela/derecho de peticion/"
                "pago oficioso (no por radicado), y se emparejara cada correo con su proceso al final.",
                credenciales_correo[0],
            )

    carpetas_existentes = _listar_carpetas_existentes()

    # Antes de crear nada: los autos que corridas anteriores dejaron en la
    # carpeta propia de cada proceso se recogen en la carpeta comun
    # CARPETA_TERMINADOS_POR_AUTO, renombrados con el numero del proceso.
    mover_autos_de_carpetas_antiguas(terminados, carpetas_existentes)

    # Se crean TODAS las carpetas de una vez, antes de ponerse a buscar
    # contenido -- para que las 1223 queden visibles en el disco desde
    # el principio, en vez de ir apareciendo intercaladas a medida que
    # el script busca documentos (que puede tardar horas para los
    # procesos con radicado antes de siquiera llegar a los terminados).
    crear_todas_las_carpetas(con_radicado, carpetas_existentes, "Carpetas de activos/suspendidos/reorganizacion/remitida/etc")
    crear_todas_las_carpetas(terminados, carpetas_existentes, "Carpetas de terminados (pago/auto/contrato/no inicio)")

    # Procesos que ya se revisaron por completo en una corrida anterior
    # (ver ARCHIVO_PROCESADOS) -- se omiten esta vez, ni Drive ni Gmail,
    # para no repetir horas de trabajo ya hecho. Se calcula UNA vez al
    # principio (no se actualiza a mitad de esta corrida) para que el
    # paso de Gmail siga cubriendo los procesos que se acaban de revisar
    # en Drive en esta MISMA corrida.
    ya_procesados = _cargar_procesados_anteriormente()
    if ya_procesados:
        logging.info(
            "[Retomar] %d carpeta(s) ya se revisaron en una corrida anterior -- se omiten esta vez "
            "(borra %s si quieres que se revisen todas de nuevo).",
            len(ya_procesados), ARCHIVO_PROCESADOS,
        )

    con_radicado_pendiente = [p for p in con_radicado if p["nombre_carpeta"] not in ya_procesados]
    terminados_pendiente = [p for p in terminados if p["nombre_carpeta"] not in ya_procesados]

    def _tarea_con_radicado(proceso):
        try:
            procesar_con_radicado(proceso)
            _marcar_como_procesado(proceso["nombre_carpeta"])
        except Exception as error:
            logging.error("[Error] Proceso %s (fila(s) %s) fallo y se omite -- se sigue con el resto: %s", proceso["numero"], ", ".join(str(f) for f in proceso["filas_excel"]), error)

    if con_radicado_pendiente:
        logging.info(
            "Buscando en Drive %d proceso(s) (hasta %d a la vez, ver NUM_HILOS)...",
            len(con_radicado_pendiente), NUM_HILOS,
        )
        with ThreadPoolExecutor(max_workers=NUM_HILOS) as pool:
            list(pool.map(_tarea_con_radicado, con_radicado_pendiente))

    if credenciales_correo:
        try:
            correos = buscar_correo_global_informacion_no_procesal(*credenciales_correo)
            logging.info(
                "[Correo] %d correo(s) de tutela/derecho de peticion/pago oficioso encontrados en todo Gmail -- "
                "emparejando con los procesos...", len(correos),
            )
            procesar_correos_no_procesal(correos, con_radicado_pendiente, carpetas_existentes)
        except Exception as error:
            logging.error("[Correo] Fallo la busqueda global en Gmail, se omite: %s", error)

    pendientes = []

    def _tarea_terminado(proceso):
        try:
            encontrado = procesar_terminado(proceso, pendientes)
            if encontrado:
                _marcar_como_procesado(proceso["nombre_carpeta"])
        except Exception as error:
            logging.error("[Error] Proceso %s (fila(s) %s) fallo y se omite -- se sigue con el resto: %s", proceso["numero"], ", ".join(str(f) for f in proceso["filas_excel"]), error)

    if terminados_pendiente:
        logging.info(
            "Buscando en Drive %d proceso(s) terminado(s) (hasta %d a la vez, ver NUM_HILOS)...",
            len(terminados_pendiente), NUM_HILOS,
        )
        with ThreadPoolExecutor(max_workers=NUM_HILOS) as pool:
            list(pool.map(_tarea_terminado, terminados_pendiente))

    if pendientes:
        with open(ARCHIVO_PENDIENTES_TERMINADOS, "w", newline="", encoding="utf-8-sig") as f:
            escritor = csv.writer(f, delimiter=";")
            escritor.writerow(["No.", "Filas Excel", "Estado", "Radicado", "Cuentas", "Demandados", "Juzgado"])
            for proceso in pendientes:
                escritor.writerow([
                    proceso["numero"], ", ".join(str(f) for f in proceso["filas_excel"]), proceso["estado"],
                    proceso["radicado"] or "", ", ".join(proceso["cuentas"]), ", ".join(proceso["demandados"]),
                    proceso["juzgado"],
                ])
        logging.warning(
            "%d proceso(s) terminado(s) se quedaron SIN el documento que los termina -- descargalos a mano, "
            "quedaron listados en %s.", len(pendientes), ARCHIVO_PENDIENTES_TERMINADOS,
        )

    if sin_estado:
        logging.warning("%d fila(s) no tienen ESTADO PROCESAL diligenciado todavia en el Excel:", len(sin_estado))
        for proceso in sin_estado:
            logging.warning("   - Fila %s, proceso %s", proceso["fila_excel"], proceso["numero"])

    if MODO_PRUEBA:
        logging.info(
            "MODO_PRUEBA esta activo: no se creo ninguna carpeta ni se descargo nada de verdad todavia. "
            "Revisa el log y, si se ve bien, cambia MODO_PRUEBA = False al inicio de este script y vuelve a correrlo."
        )


def main():
    configurar_logging()
    procesar()


if __name__ == "__main__":
    main()
