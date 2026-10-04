# Versión 4: reescritura de estilo sobre la v3. Quita relleno, explicaciones repetidas,
# advertencias duplicadas, aperturas previsibles y cierres obvios; une ideas cortadas en
# párrafos con desarrollo, y pasa los perfiles individuales (que estaban dos veces en prosa)
# a una tabla APA. No cambia datos, citas ni el mensaje del trabajo.
# Uso: python version4.py "Trabajo lectoescritura v3.docx" "Trabajo lectoescritura v4.docx"
import sys
from copy import deepcopy

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

import reescribir as R

GG = "como se citó en Giraldo Gaviria & Caro Lopera, 2022"
CT = "como se citó en Caamaño Tomás, 2021"

# índice en la v3 -> (comienzo esperado del texto actual, [párrafos nuevos])
CAMBIOS = {
    # ------------------------------------------------ Dedicatoria y agradecimientos
    26: ("A mi esposo. Te fuiste", ["A mi esposo, que, aunque ya no está, sigue acompañando cada uno de mis pasos. Este trabajo también es tuyo, porque en él quedó algo de lo que vivimos juntos y porque en los días más difíciles fue tu recuerdo el que me empujó a no rendirme."]),
    27: ("A mi mamá, que nunca", ["A mi mamá, que nunca dejó de estar, por todo lo que se esforzó para que yo pudiera estudiar y por enseñarme con su ejemplo a ser responsable y a no soltar las cosas cuando se ponen duras."]),
    28: ("Este trabajo no es solo", ["Detrás de este logro hay sacrificios de ustedes dos que casi nunca se vieron y una confianza en mí que nunca me faltó. A ambos, con todo mi cariño, se lo dedico."]),
    29: ("A los dos", []),
    39: ("Gracias a Dios", ["Agradezco a Dios por darme fuerzas para terminar este proceso, aun en los momentos en que todo se puso cuesta arriba."]),
    40: ("A mi mamá, otra vez", ["A mi mamá, porque sin sus palabras de ánimo, su paciencia y su manera de estar siempre ahí, este trabajo no existiría; es tan suyo como mío."]),
    41: ("Al profesor Jean Paul", ["Al profesor Jean Paul, por acompañarme, orientarme y confiar en este proyecto. Su ayuda, en lo académico y en lo humano, me sirvió para ordenar mis ideas y llevarlo a término."]),
    42: ("Y a mi esposo", ["Y a mi esposo, que sigue siendo parte de mi vida y de lo que consigo. Pensar en él me dio fuerza en los peores momentos, y este logro también honra su memoria."]),
    43: ("Gracias a todos", []),

    # ------------------------------------------------ Resumen y Abstract (APA: un solo párrafo)
    100: ("En este trabajo revisamos", ["Este trabajo examina los factores neurobiológicos, pedagógicos, socioculturales y tecnológicos que intervienen en la enseñanza y el aprendizaje de la lectoescritura en América Latina, con atención especial a la dislexia, la formación docente, la desigualdad educativa y las secuelas de la pandemia. Combina una revisión documental de investigaciones publicadas entre 2018 y 2024 con una prueba diagnóstica de pseudopalabras, diseñada a partir del PROLEC-R y aplicada a diez estudiantes de básica para observar la ruta fonológica, la decodificación y la conciencia fonológica. La revisión indica que la dislexia sigue siendo poco comprendida en la escuela, que la lateralidad cruzada y una conciencia fonológica débil dificultan detectar a tiempo los problemas lectores, que muchos docentes no han sido formados para atenderlos y que la pandemia amplió las brechas en los contextos vulnerables. En la prueba, nueve de los diez estudiantes alteraron al menos una pseudopalabra al leer y todos lo hicieron al escribir al dictado; aislar un fonema les resultó más difícil que segmentar en sílabas o reconocer terminaciones semejantes, tareas que los diez resolvieron bien. Los estudios revisados reportan buenos resultados con intervenciones multisensoriales, lúdicas e inclusivas, y el conjunto del trabajo apunta a la necesidad de diagnosticar temprano, adaptar las estrategias a cada contexto y formar mejor a los docentes."]),
    101: ("La revisión muestra", []),
    102: ("Los estudios revisados", []),
    107: ("This study reviews", ["This study examines the neurobiological, pedagogical, sociocultural and technological factors involved in teaching and learning to read and write in Latin America, with particular attention to dyslexia, teacher training, educational inequality and the aftermath of the pandemic. It combines a documentary review of research published between 2018 and 2024 with a pseudoword diagnostic test, designed on the basis of the PROLEC-R and administered to ten primary school students to observe the phonological route, decoding and phonological awareness. The review indicates that dyslexia is still poorly understood in schools, that crossed laterality and weak phonological awareness make it harder to detect reading problems early, that many teachers have not been trained to address them, and that the pandemic widened the gaps in vulnerable settings. In the test, nine of the ten students altered at least one pseudoword when reading and all of them did so when writing from dictation; isolating a phoneme proved harder than splitting a word into syllables or recognizing similar endings, tasks that all ten solved correctly. The studies reviewed report good results with multisensory, play-based and inclusive interventions, and the work as a whole points to the need for early diagnosis, strategies adapted to each context and better teacher preparation."]),
    108: ("The review shows", []),
    109: ("The studies reviewed", []),

    # ------------------------------------------------ Introducción
    122: ("Aprender a leer y a escribir abre", ["Aprender a leer y a escribir condiciona casi todo lo que un niño hace después en la escuela, desde comprender un enunciado hasta participar en una cultura que se organiza por escrito. Aun así, buena parte de los estudiantes latinoamericanos arrastra dificultades persistentes en ambos procesos, y la investigación las asocia con causas de distinto orden, desde la dislexia y una formación docente insuficiente hasta la inequidad que la pandemia hizo todavía más visible."]),
    123: ("Lo que queremos entender", ["Este trabajo busca entender cómo se combinan esas variables en la manera en que los estudiantes leen y escriben, y qué consecuencias tiene esa combinación para planear la enseñanza. Para ello reúne la teoría y las investigaciones recientes con lo observado al aplicar, como ejercicio piloto, una prueba diagnóstica de pseudopalabras, y se guía por cuatro preguntas:"]),
    124: ("Las preguntas que guiaron", []),
    129: ("Este estudio se justifica", ["La razón de fondo es práctica. En muchas aulas las dificultades de lectura y escritura se detectan tarde y se atienden con estrategias que no consideran las necesidades concretas de cada estudiante, y ofrecer referentes para diagnosticarlas a tiempo tiene consecuencias directas sobre el rendimiento académico, la inclusión y la participación social de los niños."]),

    # ------------------------------------------------ Planteamiento del problema
    160: ("Leer y escribir son la base", ["En muchos contextos educativos, y de manera especial en América Latina, persisten dificultades de lectura y escritura que afectan el desempeño académico de los estudiantes. Se manifiestan en errores repetidos al decodificar, en una conciencia fonológica débil, en problemas de comprensión y en una producción escrita limitada, y obligan a preguntarse por sus causas y por la manera de atenderlas desde la pedagogía y el contexto de cada estudiante."]),
    161: ("Investigaciones recientes apuntan", ["Las investigaciones recientes señalan varios factores que agravan el problema, entre ellos la dislexia, la lateralidad cruzada, las brechas socioeconómicas, la falta de formación docente y los efectos de la pandemia. Muchos maestros, además, no cuentan con herramientas para detectar a tiempo las fallas en la ruta fonológica ni para intervenir sobre ellas, de modo que la distancia entre lo que se espera que un estudiante lea y escriba y lo que efectivamente logra tiende a ampliarse."]),
    162: ("El problema que nos ocupa", ["El problema de investigación consiste en comprender cómo interactúan los factores neurocognitivos, pedagógicos y socioculturales en la aparición de estas dificultades, y en establecer hasta qué punto una prueba de pseudopalabras permite identificar fallas específicas en la decodificación y en la conciencia fonológica que sirvan de base para diseñar estrategias de enseñanza y acompañamiento."]),

    # ------------------------------------------------ Antecedentes
    179: ("Varias investigaciones coinciden", ["Las investigaciones revisadas coinciden en que el aprendizaje lector depende tanto de la enseñanza como de procesos neurocognitivos como la memoria de trabajo, la velocidad de procesamiento, la conciencia fonológica y las rutas de acceso al léxico. En los estudios sobre dislexia, las dificultades para decodificar y reconocer patrones fonológicos aparecen incluso en estudiantes sin diagnóstico previo, lo que respalda la insistencia en evaluar temprano y con instrumentos adecuados."]),
    180: ("Los trabajos revisados también", ["Varios trabajos relacionan además la lateralidad cruzada y la dislexia con omisiones, sustituciones e inversiones de letras y con dificultades para automatizar la lectura, mientras que otros documentan el golpe que la pandemia dio a la alfabetización en los contextos vulnerables, donde la falta de recursos tecnológicos y pedagógicos profundizó las brechas. Los estudios sobre formación docente, por su parte, muestran una distancia considerable entre lo que exige el aula y lo que los maestros saben sobre alfabetización."]),
    181: ("En cuanto a la formación docente", []),

    183: ("Este artículo trata la dislexia", ["Aguirre-Medrano y González-López (2021) estudian la relación entre dislexia y lateralidad en la escritura a partir de un caso de educación básica, y en particular cómo una lateralidad cruzada lleva a confundir grafemas como b-d y p-q. El diagnóstico combina el PROLEC-R, el WISC IV, el Test de Harris de lateralidad y pruebas de discriminación fonológica."]),
    184: ("En el caso estudiado", ["La niña evaluada tiene una lateralidad cruzada mal afirmada, un rasgo que llama la atención porque, según Duarte-Hernández y Pérez-Mendoza (2020), quienes evaluaron con el test de Harris a niños de 2 a 5 años, la lateralidad suele quedar definida hacia los 6 años. Presenta además disgrafía fonológica y dificultades para identificar letras y signos de puntuación, y aunque su coeficiente intelectual supera el promedio (CI = 115), lee y escribe con lentitud. Las autoras concluyen que la lateralidad pesa en el desarrollo lector de los niños con dislexia más de lo que la escuela suele reconocer, y recomiendan intervenciones psicopedagógicas que trabajen la conciencia fonológica, la discriminación auditiva y la visomotricidad."]),
    185: ("El estudio recomienda", ["El caso está descrito con detalle, pero al tratarse de una sola paciente sus hallazgos no se pueden generalizar, y el estudio no incluye un seguimiento que permita saber si las intervenciones funcionaron con el tiempo, algo relevante porque la conciencia fonológica y la escritura se desarrollan por etapas (Gutiérrez Fresneda & Díez Mediavilla, 2018). Tampoco considera la influencia de la escuela y la familia, que puede ser decisiva en la evolución de la dislexia (Snowling et al., 2020). Para el docente, su principal aporte es mostrar por qué conviene reforzar la conciencia fonológica y las habilidades visomotoras de manera individualizada."]),
    186: ("Desde una mirada crítica", []),
    187: ("En la práctica, el artículo", []),

    189: ("Giraldo Gaviria y Caro Lopera (2022) analizan", [f"Giraldo Gaviria y Caro Lopera (2022) parten de que la escritura académica es indispensable en la formación de maestros, porque les da acceso a las comunidades discursivas de su disciplina y desarrolla sus competencias para investigar y argumentar. Para sustentarlo hacen una revisión sistemática exploratoria, apoyada en los planteamientos de Manchado (2009, {GG}), de 60 textos académicos publicados en el mundo hispanohablante durante los últimos 17 años, que organizan en tres categorías:"]),
    190: ("Estos autores realizan", []),
    191: ("1. Alfabetización académica", [f"1. Alfabetización académica en la formación de nuevos maestros. La escritura académica aparece como condición para adquirir conocimientos y desarrollar pensamiento crítico, y para Carlino (2005) y Vargas Franco (2020, {GG}) es, además, un proceso de construcción de identidad académica."]),
    192: ("2. Escritura académica", [f"2. Escritura académica en la formación docente. Molina (2017, {GG}) encuentra que la mayoría de los estudiantes ve la escritura como un medio para certificar lo que sabe y no como una herramienta para pensar. Carlino et al. (2013, {GG}), por su parte, sostienen que debería enseñarse dentro de todas las disciplinas."]),
    193: ("3. Escritura, investigación", [f"3. Escritura, investigación y formación en las ENS. Zambrano y Aragón de Moreno (2015, {GG}) proponen cambios curriculares y metodológicos para que los programas de formación complementaria enseñen a escribir con propósito investigativo."]),
    194: ("Desde un punto de vista crítico", ["La revisión defiende con buenos argumentos que la escritura académica sea un eje de la formación inicial docente, entendida, según los autores, como un proceso social y epistémico que da acceso a la cultura escrita de las disciplinas. Lo que no explica es cómo llevar esos cambios a las ENS colombianas de hoy, y esa ausencia limita el alcance de su propuesta."]),
    195: ("Para concluir, la alfabetización", []),

    197: ("El autor plantea cómo", [f"Caamaño Tomás (2021) sostiene que la pandemia de COVID-19 transformó la educación al trasladar la enseñanza, y con ella la lectura y la escritura, a entornos digitales, y describe las barreras tecnológicas y pedagógicas de ese tránsito. Para explicarlo recurre a La galaxia Gutenberg: génesis del Homo Typographicus, donde Marshall McLuhan (1962, {CT}) distingue tres momentos en la transmisión del conocimiento: el alfabeto fonético, la imprenta y la era electrónica. A su juicio, la adaptación a la educación digital depende menos de los planes de estudio o de las políticas que de la preparación técnica de estudiantes y docentes."]),
    198: ("Según el ensayo", []),
    199: ("Caamaño Tomás lee el impacto", [f"Según McLuhan, los entornos electrónicos reintroducen lo auditivo y lo visual y por eso recuperan rasgos de la oralidad anteriores a la escritura fonética (McLuhan, 1962, {CT}). Caamaño Tomás matiza esa lectura con lo que se perdió en la virtualidad, como los gestos y la interacción en el aula, y con la individualización que produjeron el encierro y la dependencia de las clases en línea, que debilitó la relación entre docentes y alumnos y con ella la colaboración. De ahí que, retomando a McLuhan (1969, {CT}), proponga repensar la educación a partir de la manera en que los medios cambian la percepción y la transmisión de la cultura."]),
    200: ("El ensayo también se fija", []),
    201: ("Al cierre, Caamaño Tomás", [f"El ensayo cierra con la advertencia de Sven Birkerts (1994, {CT}), autor de Elegía a Gutenberg, sobre la capacidad de la tecnología para transformar la conciencia humana. No ofrece soluciones, pero deja planteada la necesidad de formar lectores y escritores capaces de usar los entornos digitales con sentido crítico."]),

    203: ("El estudio de Pisco-Román", ["Pisco-Román y Bailón-Panta (2023) analizan, con un enfoque cuantitativo y descriptivo, la lectoescritura de estudiantes de Básica Media de una institución educativa de Manabí, en Ecuador. A partir de encuestas, observación y revisión documental encuentran que la comprensión lectora es una de las principales dificultades del grupo y que afecta buena parte de su desempeño académico, por lo que recomiendan fortalecer la lectoescritura desde los primeros años con estrategias innovadoras."]),
    204: ("Los autores parten de que", ["Los autores entienden la lectura y la escritura como procesos cognitivos y sociales complejos que se desarrollan al mismo tiempo (Pisco-Román & Bailón-Panta, 2023, p. 331) y que sostienen el aprendizaje escolar, porque en ellos se forman habilidades cognitivas, comunicativas y lingüísticas (p. 333). Quien lee y escribe bien, añaden, puede organizar sus ideas, comprender textos complejos, pensar críticamente y expresarse con claridad (p. 335)."]),
    205: ("Por eso consideran", []),
    206: ("Según ellos, quien lee", []),
    207: ("El artículo tiene una estructura", ["El artículo está bien organizado, aunque repite algunas ideas y deja apartados poco desarrollados, y su título promete más de lo que cumple, porque el texto termina hablando del proceso de enseñanza-aprendizaje en general más que de la lectoescritura."]),

    209: ("Morales Londoño (2018) estudia", ["Morales Londoño (2018) revisa la calidad de la enseñanza de la lectoescritura en la básica primaria colombiana a partir de los resultados de las pruebas SABER y PISA, los lineamientos del Ministerio de Educación Nacional (MEN) y las estrategias gubernamentales para mejorarla. Para definir la calidad retoma a Nyathi et al. (2011, como se citó en Morales Londoño, 2018), que la entienden como un concepto multidimensional, multinivel y dinámico, ligado al contexto del modelo educativo, a la misión y los objetivos de cada institución y a los estándares de cada programa, disciplina o sistema."]),
    210: ("Para hablar de calidad", []),
    211: ("El artículo se estructura", ["Al tratarse solo de una revisión documental, el trabajo no contrasta sus afirmaciones con datos propios, no profundiza en el papel de los docentes ni evalúa el efecto real de las estrategias que analiza, y presenta además algunos problemas de cohesión."]),

    213: ("El estudio de Rivera Cintrón", ["Rivera Cintrón y Batiz Cartagena (2024) proponen enseñar la lectoescritura en una escuela pública de Nueva York con el currículo Artes del Lenguaje, basado en el neuroaprendizaje y en estrategias lúdicas. Retoman a Vargas et al. (2019, como se citó en Rivera Cintrón & Batiz Cartagena, 2024, p. 8), para quienes leer y escribir son una muestra de conectividad intelectual y neuronal por tratarse de uno de los aprendizajes más complejos que realizan las personas. Desde ahí defienden el juego como el camino más agradable para que un niño aprenda (p. 9)."]),
    214: ("Los autores retoman a Vargas", []),
    215: ("Entre sus hallazgos", []),
    216: ("Aunque la propuesta resulta", ["La propuesta mejoró la motivación de los estudiantes, pero la muestra fue pequeña, hubo problemas administrativos y no se midió con pruebas estandarizadas cuánto avanzó la lectoescritura, de modo que sus conclusiones necesitan una evaluación más amplia."]),

    218: ("Tinta Aruquipa (2020) aborda", ["Tinta Aruquipa (2020) aborda la escritura desde la percepción sensorial y la lectura de la realidad, y encuentra un alto porcentaje de estudiantes con dificultades para escribir, frente a las cuales destaca el papel de la motivación y del enfoque sensorial."]),
    219: ("Sin embargo, la exposición", ["La exposición teórica es dispersa, ya que menciona a Vygotsky sin articularlo con la propuesta, describe con poca precisión los instrumentos y no indaga en las causas de las dificultades ni propone estrategias concretas para atenderlas."]),

    221: ("López Rivas (2024) analiza", ["López Rivas (2024) estudia, a partir de casos urbanos y rurales de Guatemala, cómo se enseñó a leer y escribir durante la pandemia, cuando el cierre de las escuelas desordenó las clases y golpeó con especial fuerza a la lectoescritura, que exige un acompañamiento cercano del docente (p. 2). El paso a la modalidad a distancia obligó a los maestros a inventar estrategias para llegar a sus estudiantes (p. 3), y el autor destaca la resiliencia y la creatividad con que mantuvieron el vínculo en medio de las desigualdades, la falta de acceso a la tecnología y la dependencia del apoyo de las familias (p. 3)."]),
    222: ("Según el autor, la pandemia", []),
    223: ("Las escuelas tuvieron", []),
    224: ("Uno de los hallazgos principales", []),
    225: ("Aunque el estudio aporta", ["La muestra es limitada y el estudio no aporta datos sobre la mejora de la lectoescritura, por lo que sus hallazgos ganarían si se compararan con los de otros contextos."]),

    227: ("Domínguez Vázquez (2023) presenta", ["Domínguez Vázquez (2023) diseña y aplica una intervención psicopedagógica para mejorar la lectoescritura de estudiantes de segundo de primaria en riesgo de exclusión social, con un enfoque de investigación-acción y un diseño cuasiexperimental, es decir, sin asignación aleatoria de los participantes (Arnau, 1995, como se citó en Núñez Peña, 2011). El programa, enmarcado en el Objetivo de Desarrollo Sostenible de educación inclusiva y de calidad (Naciones Unidas, 2015), combina lectura interactiva, reconocimiento de fonemas y escritura creativa con actividades multisensoriales y lúdicas, y mide de forma cualitativa y cuantitativa el progreso en comprensión lectora, escritura, reconocimiento de palabras, conciencia fonológica y motivación."]),
    228: ("En lo teórico, el estudio", ["La autora reporta mejoras importantes en todas esas áreas, aunque reconoce que el alto ausentismo del grupo interrumpió la continuidad que exige consolidar competencias básicas, y propone por eso contar con personal especializado, como los profesores de servicios a la comunidad (PSC). Medir a la vez lo cognitivo y lo afectivo es uno de los aciertos del trabajo. Su principal límite es metodológico, ya que la ausencia de un grupo de control aleatorizado y el hecho de que todo ocurriera en una sola escuela impiden extender los resultados a otras poblaciones, aunque el diseño flexible permitió ajustar la intervención sobre la marcha."]),
    229: ("Lo más interesante es cómo", []),
    230: ("El artículo no oculta", []),
    231: ("La investigación mezcla", []),
    232: ("El estudio tiene límites", []),
    233: ("En suma, Domínguez Vázquez", []),

    # ------------------------------------------------ Justificación
    236: ("Leer y escribir bien es indispensable", ["Las dificultades de lectura y escritura siguen apareciendo en las escuelas, sobre todo donde se cruzan la dislexia, la lateralidad cruzada, las brechas socioculturales y una formación docente insuficiente, y en las instituciones suelen detectarse tarde y atenderse con estrategias que rara vez consideran el contexto de cada grupo. Como la lectura de pseudopalabras permite observar la ruta fonológica y la decodificación con una precisión que no ofrecen las palabras conocidas, decidimos diseñar y analizar una prueba diagnóstica de este tipo."]),
    237: ("Escogimos este tema", ["El trabajo reúne en una misma revisión perspectivas neurocognitivas, pedagógicas y socioculturales que suelen estudiarse por separado, y deja un instrumento y un análisis que pueden servir a otras investigaciones, a programas de intervención y a la formación de docentes interesados en atender estas dificultades a tiempo."]),
    238: ("El trabajo reúne en una misma", []),

    # ------------------------------------------------ Marco teórico
    255: ("Entendemos la lectoescritura", ["El marco teórico combina dos perspectivas. La primera entiende la lectoescritura como un proceso cognitivo en el que el lector decodifica, comprende, retiene lo que va leyendo y construye significado, y en el que la conciencia fonológica cumple un papel central; la segunda, pedagógica y contextual, recoge lo que la investigación dice sobre el peso del entorno, de la pandemia y de la formación docente. Ambas confluyen en la prueba de pseudopalabras, que permite observar la ruta fonológica de cada estudiante y ofrece así un punto de partida para diseñar estrategias acordes con sus necesidades y su contexto."]),
    256: ("Para observar esa ruta", []),
    257: ("A esa mirada cognitiva", []),
    258: ("Con estas dos miradas", []),
    262: ("La lectoescritura reúne habilidades", ["La lectoescritura reúne habilidades cognitivas, lingüísticas y motoras necesarias para interpretar, producir y comprender textos, y exige decodificar, sostener la atención, usar la memoria de trabajo y construir significado. Por esa amplitud, Morales Londoño (2018) la considera una base del desempeño académico y de la participación social, y dedica buena parte de su análisis a las ideas equivocadas que circulan sobre cómo se aprende y sobre el papel del docente."]),
    263: ("Morales Londoño (2018) revisa", []),
    265: ("La conciencia fonológica es la capacidad", ["La conciencia fonológica es la capacidad de identificar y manipular los sonidos del habla, y de ella depende en buena parte que el niño logre relacionar de manera estable grafemas y fonemas. Su papel se hace más visible en estudiantes con dislexia, disgrafía fonológica o problemas de discriminación auditiva, en quienes Aguirre-Medrano y González-López (2021) documentan los errores típicos de la ruta fonológica que detectan pruebas como el PROLEC-R (omisiones, sustituciones, inversiones y segmentaciones incorrectas), además de confusiones entre grafemas como b-d y p-q cuando la lateralidad está mal afirmada."]),
    266: ("Pruebas como el PROLEC-R", []),
    267: ("Las pseudopalabras, que forman", []),
    270: ("Los estudios de Domínguez", ["Domínguez Vázquez (2023) muestra que, en contextos de riesgo social, las intervenciones diseñadas para las necesidades del grupo mejoran la lectoescritura, y Pisco-Román y Bailón-Panta (2023) identifican la comprensión lectora como una de las dificultades más comunes de la básica media, con efectos sobre el rendimiento que exigen estrategias más innovadoras."]),
    271: ("López Rivas (2024), por su parte", []),
    273: ("La calidad de la enseñanza", ["La calidad de esa enseñanza depende también de cómo se forma a los maestros. Giraldo Gaviria y Caro Lopera (2022) piden fortalecer la alfabetización académica desde la formación inicial, y Morales Londoño (2018) atribuye las fallas de la enseñanza de la lectoescritura en Colombia a concepciones equivocadas, prácticas tradicionales y poca actualización profesional."]),
    275: ("Aprender a leer no depende", ["El aprendizaje lector depende también del contexto familiar, de las oportunidades de acceder a textos, de las metodologías docentes y de las condiciones socioculturales. López Rivas (2024) muestra que durante la pandemia el acceso limitado a la tecnología y el nivel educativo de los padres afectaron de manera distinta la alfabetización en zonas urbanas y rurales. Los estudiantes con menos recursos tuvieron más dificultades para sostener su proceso, y de ahí la necesidad de programas de apoyo y seguimiento escolar y de metodologías flexibles."]),
    276: ("Esto dejó en evidencia", []),
    278: ("Las pruebas de pseudopalabras sirven", ["Una pseudopalabra como “trunal” no pertenece al léxico del lector, que por eso no puede reconocerla de memoria y debe leerla aplicando las reglas de conversión grafema-fonema. Esa característica convierte las pruebas de pseudopalabras en un recurso adecuado para evaluar la ruta fonológica y explica que el PROLEC-R, en cuya tarea de pseudopalabras nos basamos para diseñar la prueba, las incluya. Al eliminar el apoyo de la memorización, dejan ver el dominio real de las reglas fonológicas y grafémicas, tanto al leer como al escribir."]),
    279: ("En la misma línea, Aguirre", ["Los errores que aparecen en estas tareas son, además, los que la literatura asocia con dificultades de decodificación y de conciencia fonológica (Aguirre-Medrano & González-López, 2021), y su análisis ofrece pistas concretas para diseñar intervenciones individualizadas sobre la decodificación, la discriminación fonológica y el reconocimiento de las unidades gráficas."]),
    280: ("La batería PROLEC-R, en la que", []),
    281: ("Analizar los errores", []),

    # ------------------------------------------------ Metodología
    297: ("Como el estudio buscaba", ["La investigación tiene un enfoque cualitativo, porque busca comprender cómo lee y escribe cada uno de los diez estudiantes y qué tipo de errores comete, e interpretar esos resultados a la luz de sus condiciones particulares."]),
    299: ("El trabajo es un estudio de casos", ["Se trata de un estudio de casos múltiples de tipo descriptivo, en el que el desempeño de cada estudiante se analiza por separado y luego se compara con el de los demás. El diseño no pretende generalizar, sino obtener información detallada sobre procesos concretos de lectura y escritura, como corresponde al carácter diagnóstico del trabajo."]),
    301: ("Participaron diez estudiantes", ["Participaron diez estudiantes de educación básica, evaluados de manera individual e identificados con números del 1 al 10, tal como aparecen en la transcripción (Anexo A)."]),
    303: ("El alcance es descriptivo", ["El alcance es descriptivo. Se caracterizan las dificultades de lectura y escritura a partir de los errores y respuestas registrados en la prueba, sin pretender explicar sus causas ni establecer relaciones estadísticas."]),
    305: ("El método fue evaluativo", ["El método, evaluativo-descriptivo, tuvo dos partes. La primera fue una revisión documental de artículos publicados entre 2018 y 2024 en América Latina, España y Estados Unidos sobre dislexia, lateralidad, formación docente, pandemia y estrategias de enseñanza de la lectoescritura, cuyos resultados se presentan en los antecedentes y en el marco teórico. La segunda consistió en aplicar una prueba diagnóstica de pseudopalabras, diseñada para este estudio a partir de la tarea de pseudopalabras del PROLEC-R, y en analizar las respuestas con los criterios de la propia prueba y con los referentes teóricos del trabajo."]),
    307: ("La técnica de recolección", ["La técnica fue la evaluación individual. Cada estudiante resolvió las cinco actividades de la prueba, cuyas respuestas se registraron por escrito (Anexo A) y se compararon después, una por una, con los estímulos originales:"]),
    309: ("Las respuestas de cada estudiante", []),
    310: ("Usamos pseudopalabras porque", []),

    # ------------------------------------------------ Resultados
    327: ("Los resultados de la prueba de pseudopalabras nos dan", []),
    328: ("Trabajar con pseudopalabras", []),
    329: ("Lectura de pseudopalabras:", ["Lectura de pseudopalabras. Nueve de los diez estudiantes alteraron al menos una de las treinta pseudopalabras de la lista, aunque en cantidades muy distintas, desde ninguna en el estudiante 4 hasta nueve en los estudiantes 7 y 8. Las modificaciones más frecuentes fueron “Zepa”, “Zeman” o “Zapa” por “Zepan”, “Blused” por “Lused”, “Florir” por “Folir” y “Flanit” por “Franit”, junto con otras como “Trampin”, “Digan” o “Zolil”, y en casi todos los casos consistieron en cambiar, añadir u omitir una o dos letras sin perder la estructura de la pseudopalabra."]),
    330: ("Esto importa para analizar", []),
    331: ("Escritura al dictado de pseudopalabras:", ["Escritura al dictado. Los diez estudiantes escribieron alguna pseudopalabra de forma distinta al estímulo, con producciones como “Lisa” por “Plisa”, “Bieco” por “Bieclo”, “Grac” por “Gruac”, “Tranflopiar” por “Tramkopliar”, “Caudre” por “Cuadre”, “Fragno” por “Fradno” o “Pamil” por “Pamir”. A los estudiantes 1 y 2 se les dictó una lista y a los demás otra, de modo que las producciones de los dos grupos no son comparables palabra por palabra."]),
    332: ("Las dificultades, entonces", []),
    333: ("Sin embargo, estas modificaciones", []),
    334: ("Comprensión y conciencia fonológica:", ["Conciencia fonológica. El sonido inicial de “gropel” lo aislaron cinco estudiantes, mientras los otros cinco respondieron “Gro”. El final de “finod” solo lo identificaron cuatro, y entre los demás aparecieron “Nod”, “Not”, “Od” y hasta “F”, que es el sonido inicial; con el interno de “lumep” hubo también cuatro aciertos, frente a respuestas como “Met”, “Ume”, “Me” o “Um”. Las otras dos tareas no presentaron dificultad, porque los diez dividieron “trunal” como “Tru-nal” y todos señalaron “zopar” y “perar” como las pseudopalabras que terminan igual."]),
    335: ("Para el sonido inicial", []),
    336: ("El desempeño cambia, pues", []),
    337: ("En contraste, la segmentación", []),
    338: ("Estos resultados permiten establecer", []),
    339: ("Fluidez lectora en frases:", ["Fluidez lectora en frases. Ningún estudiante alteró las pseudopalabras incluidas en las tres oraciones (“El drumo saltó sobre la mesa”, “María compró un teniq azul” y “El perro jugaba con el bapo en el parque”). Los cambios se concentraron en los artículos, ya que ocho estudiantes dijeron “un bapo” en lugar de “el bapo” y el estudiante 1 omitió además el artículo inicial de la primera frase y cambió “El perro” por “Mi perro”."]),
    340: ("El dato es interesante", []),
    341: ("Asociación grafema–fonema:", ["Asociación grafema-fonema. La mitad de los estudiantes (1, 2, 5, 8 y 10) identificó /Z/ como sonido inicial de “zunel” y la otra mitad respondió con la sílaba “Zu”. Con “quarim” hubo más dispersión, pues solo los estudiantes 8 y 10 señalaron “q”, seis respondieron con segmentos como “qu”, “qua” o “cua” y los estudiantes 7 y 9 dieron “au” y “ua”. La última parte de la actividad, unir pseudopalabras con su sonido inicial, no quedó registrada en la transcripción y no se analizó."]),
    342: ("Hay, entonces, distintos niveles", []),
    343: ("Diferencias individuales.", ["Desempeño por estudiante. La Tabla 1 reúne los resultados de cada participante."]),
    344: ("Los registros también muestran", ["Los perfiles no siguen un patrón único. El estudiante 4 leyó toda la lista sin errores y acertó todas las preguntas fonológicas, pero tuvo transformaciones al escribir. El 10 cambió cuatro pseudopalabras al leer y aun así resolvió con precisión todas las tareas de conciencia fonológica y de asociación grafema-fonema, mientras que el 3 apenas alteró dos pseudopalabras y fue de los que más dificultad tuvo para aislar sonidos."]),
    345: ("El estudiante 3 cambió solo", []),
    346: ("En los estudiantes 6 a 10", []),
    347: ("En resumen, encontramos", []),
    348: ("Las dificultades más claras", []),

    # ------------------------------------------------ Análisis de resultados
    354: ("Al analizar la transcripción", [
        "Los resultados coinciden en buena medida con lo que el marco teórico plantea sobre la ruta fonológica. Las pseudopalabras, que obligan a decodificar sin apoyo de la memoria léxica, concentraron los errores: nueve estudiantes alteraron alguna al leer, todos lo hicieron al escribir y solo tres resolvieron sin fallas las cinco tareas de conciencia fonológica. Como la mayoría de esos errores modificó una o dos letras sin impedir la lectura, la dificultad parece estar en la precisión de la conversión grafema-fonema más que en la capacidad de decodificar.",
        "La escritura al dictado plantea una exigencia distinta, porque el estudiante debe retener la secuencia sonora, distinguir sus unidades y elegir los grafemas sin tener el estímulo delante. Que las transformaciones aparecieran en los diez registros y fueran del mismo tipo que en la lectura, omisiones y sustituciones de letras como “Lisa” por “Plisa” o “Grac” por “Gruac”, sugiere que en varios estudiantes la representación fonológica todavía no es estable. Esto no permite atribuir cada error a una sola causa, pues una producción distinta del estímulo puede deberse también a la percepción auditiva, a la memoria inmediata o al conocimiento ortográfico. Además, como las listas de lectura y de dictado no eran las mismas, la relación que plantea la cuarta pregunta de investigación solo puede establecerse en el tipo de error, no en la palabra concreta.",
        "La conciencia fonológica es el componente en el que la prueba precisa más el diagnóstico. Los diez estudiantes segmentaron bien en sílabas y reconocieron las terminaciones semejantes, de modo que su problema no puede describirse como una conciencia fonológica baja en general. Las dificultades se concentran en aislar un fonema, ya que la mitad respondió “Gro” cuando se le pidió el sonido inicial de “gropel”, seis no identificaron el final de “finod” y otros seis tampoco el interno de “lumep”, casi siempre porque dieron la sílaba o un segmento más largo en lugar del sonido. El mismo fenómeno reaparece en la asociación grafema-fonema, donde cinco estudiantes respondieron “Zu” y no /Z/ para “zunel”.",
        "En las frases, en cambio, ningún estudiante alteró las pseudopalabras. Es posible que el contexto sintáctico y semántico facilite la lectura, pero la comparación con la lista tiene un límite claro, porque “drumo”, “teniq” y “bapo” eran de las pseudopalabras más cortas y sencillas, “teniq” y “bapo” se leyeron bien también de forma aislada y la transcripción no registra tiempos ni pausas. Confirmar el efecto del contexto exigiría usar las mismas pseudopalabras en ambas tareas.",
        "Ninguno de estos resultados basta para diagnosticar dislexia, que requiere una evaluación integral capaz de descartar otras explicaciones. Tampoco permiten afirmar lateralidad cruzada, porque la prueba no incluyó tareas como el Test de Harris que usaron Aguirre-Medrano y González-López (2021) ni evaluó las confusiones entre b-d y p-q que estas autoras asocian con ella. Algo parecido ocurre con los factores socioculturales y pedagógicos descritos en el marco teórico. La transcripción no informa sobre el nivel educativo de los padres, el acceso a la tecnología, las prácticas de lectura en casa ni la experiencia de cada niño durante la pandemia, así que esos factores ayudan a situar el problema, pero no pueden presentarse como causas de los errores observados. Establecer esa relación habría requerido entrevistas, observaciones de clase o cuestionarios a las familias.",
        "En relación con los objetivos, la prueba permitió observar en la práctica los fundamentos sobre conciencia fonológica, decodificación y rutas lectoras identificados en el marco teórico, y el instrumento diseñado para la ruta fonológica se aplicó a los diez participantes. El objetivo de describir los factores socioculturales, económicos y pedagógicos se resolvió desde la revisión documental, por lo que el trabajo ofrece más evidencia empírica sobre los componentes fonológicos que sobre los contextuales.",
    ]),

    # ------------------------------------------------ Discusión
    424: ("Al discutir los resultados", [
        "Los resultados son coherentes con lo que la literatura revisada plantea sobre el papel de la conciencia fonológica en el aprendizaje lector (Aguirre-Medrano & González-López, 2021), y permiten precisarlo para este grupo: las dificultades no aparecen al reconocer estructuras amplias, como sílabas o rimas, sino al aislar fonemas y convertir con exactitud grafemas en sonidos. La distinción tiene consecuencias pedagógicas, porque un estudiante que segmenta bien pero no aísla fonemas necesita un trabajo distinto del que requiere quien no reconoce las sílabas, y tratar ambos casos como “baja conciencia fonológica” llevaría a intervenciones poco ajustadas.",
        "La variedad de perfiles refuerza la idea de que la lectoescritura no es una habilidad única. Hubo estudiantes que leyeron casi sin errores y fallaron al aislar sonidos, como el 3, y otros que alteraron varias pseudopalabras pero resolvieron bien todas las tareas fonológicas, como el 2 y el 10, de modo que una evaluación centrada en un solo componente puede pasar por alto dificultades o atribuirlas mal. Esto respalda la importancia que Aguirre-Medrano y González-López (2021) dan a un diagnóstico preciso y a un trabajo ajustado a cada estudiante.",
        "Los antecedentes atribuyen buena parte de las dificultades a factores pedagógicos y contextuales, como las prácticas tradicionales, los vacíos en la formación docente y los efectos de la pandemia (Morales Londoño, 2018; López Rivas, 2024). La prueba no permite comprobarlo en estos diez casos, y esa es la principal tarea que deja abierta el trabajo, la de combinar el diagnóstico fonológico con información sobre la escuela y la familia de cada estudiante para saber cuánto de lo observado responde a cada factor.",
    ]),

    # ------------------------------------------------ Conclusiones
    441: ("El análisis de la prueba de pseudopalabras nos lleva", [
        "Los diez estudiantes evaluados muestran niveles distintos de consolidación en la decodificación, la conciencia fonológica y la correspondencia grafema-fonema. Nueve alteraron al menos una pseudopalabra al leer, con un rango que va de ningún error en el estudiante 4 a nueve en los estudiantes 7 y 8, y los diez transformaron alguna al escribir al dictado. En respuesta a la segunda pregunta de investigación, las dificultades se manifiestan sobre todo como sustituciones, omisiones y adiciones de letras en estímulos que no pueden resolverse de memoria.",
        "La dificultad más clara está en aislar fonemas. Mientras los diez estudiantes segmentaron bien en sílabas y reconocieron terminaciones semejantes, solo tres resolvieron todas las tareas de conciencia fonológica, y la mayoría de los errores consistió en dar la sílaba o un segmento más largo en lugar del sonido. Los errores de lectura y de escritura son del mismo tipo, lo que apunta a un origen común en la conversión grafema-fonema, aunque esta respuesta a la cuarta pregunta debe confirmarse en una aplicación que use los mismos estímulos en las dos tareas.",
        "En cuanto a las preguntas sobre la enseñanza, los estudios revisados coinciden en que dan mejores resultados las intervenciones que trabajan de manera explícita la conciencia fonológica y la discriminación auditiva (Aguirre-Medrano & González-López, 2021). También funcionan las que combinan actividades multisensoriales y lúdicas con estudiantes en riesgo de exclusión social (Domínguez Vázquez, 2023) y las que usan el juego para motivar (Rivera Cintrón & Batiz Cartagena, 2024). Ninguna funciona sin docentes formados para aplicarla (Giraldo Gaviria & Caro Lopera, 2022; Morales Londoño, 2018), y la experiencia de la pandemia mostró que el acompañamiento de las familias y el acceso a recursos pesan tanto como el método (López Rivas, 2024).",
        "La prueba no permite diagnosticar dislexia ni lateralidad cruzada, y los factores socioculturales y pedagógicos descritos en el marco teórico no fueron medidos, por lo que no pueden presentarse como causas de lo observado. Su aporte está en mostrar que un instrumento sencillo, aplicado de forma individual, distingue dificultades que una evaluación general confundiría, y en dejar una base para diseñar estrategias diferenciadas de conciencia fonológica, decodificación y escritura que convendría aplicar desde los primeros grados.",
    ]),
}

# párrafos que se eliminan sin reemplazo (además de los que en CAMBIOS tienen lista vacía)
ELIMINAR_RANGOS = [(355, 421), (425, 431), (442, 453)]

TABLA_TRAS = 343
TABLA = {
    "numero": "Tabla 1",
    "titulo": "Desempeño de cada estudiante en la prueba de pseudopalabras",
    "encabezado": ["Estudiante", "Pseudopalabras alteradas al leer (de 30)", "Tareas fonológicas resueltas (de 5)", "Respuestas fonológicas imprecisas", "“zunel” / “quarim”", "Lectura de frases"],
    "filas": [
        ["1", "5", "4", "“F” (final de “finod”)", "/Z/ / cua", "Omite “El”; “Mi perro”; “un bapo”"],
        ["2", "5", "5", "—", "/Z/ / qua", "“un bapo”"],
        ["3", "2", "2", "“Gro”, “Not”, “Met”", "Zu / qu", "“un bapo”"],
        ["4", "0", "5", "—", "Zu / qua", "“un bapo”"],
        ["5", "1", "4", "“Ume”", "/Z/ / qu", "“un bapo”"],
        ["6", "2", "2", "“Gro”, “Od”, “Um”", "Zu / qua rim", "“un bapo”"],
        ["7", "9", "2", "“Gro”, “Nod”, “Me”", "Zu / au", "“un bapo”"],
        ["8", "9", "2", "“Gro”, “Nod”, “Me”", "/Z/ / q", "“un bapo”"],
        ["9", "4", "2", "“Gro”, “Nod”, “Ume”", "Zu / ua", "Sin cambios"],
        ["10", "4", "5", "—", "/Z/ / q", "Sin cambios"],
    ],
    "nota": "Las tareas fonológicas son sonido inicial (“gropel”), final (“finod”) e interno (“lumep”), terminación semejante (“zopar”–“perar”) y segmentación silábica (“trunal”). Los diez estudiantes transformaron alguna pseudopalabra en la escritura al dictado. Elaboración propia a partir de la transcripción (Anexo A).",
}


def borde(tc_o_tbl, lados, tipo="single"):
    pr = tc_o_tbl.find(qn("w:tblPr")) if tc_o_tbl.tag == qn("w:tbl") else tc_o_tbl.get_or_add_tcPr()
    nombre = "w:tblBorders" if tc_o_tbl.tag == qn("w:tbl") else "w:tcBorders"
    caja = pr.find(qn(nombre))
    if caja is None:
        caja = OxmlElement(nombre)
        pr.append(caja)
    for lado in lados:
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), tipo)
        el.set(qn("w:sz"), "6" if tipo != "nil" else "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        caja.append(el)


def parrafo_tras(ancla, texto, molde, negrita=False, cursiva=False):
    nuevo = deepcopy(molde._p)
    ancla._p.addnext(nuevo) if isinstance(ancla, Paragraph) else ancla.addnext(nuevo)
    q = Paragraph(nuevo, molde._parent)
    R.reemplazar(q, texto)
    q.runs[0].bold = negrita or None
    q.runs[0].italic = cursiva or None
    q.paragraph_format.first_line_indent = 0
    return q


def insertar_tabla(d, ancla, molde):
    numero = parrafo_tras(ancla, TABLA["numero"], molde, negrita=True)
    titulo = parrafo_tras(numero, TABLA["titulo"], molde, cursiva=True)
    filas = [TABLA["encabezado"]] + TABLA["filas"]
    t = d.add_table(rows=len(filas), cols=len(filas[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, fila in enumerate(filas):
        for c, valor in enumerate(fila):
            celda = t.cell(r, c)
            p = celda.paragraphs[0]
            run = p.add_run(valor)
            run.font.size = Pt(10)
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.space_after = Pt(0)
            if r == 0:
                borde(celda._tc, ["bottom"])
    borde(t._tbl, ["top", "bottom"])
    borde(t._tbl, ["left", "right", "insideH", "insideV"], "nil")
    titulo._p.addnext(t._tbl)
    nota = parrafo_tras(t._tbl, "", molde)
    for r in list(nota.runs):
        r._r.getparent().remove(r._r)
    a = nota.add_run("Nota. ")
    a.italic = True
    nota.add_run(TABLA["nota"])
    for r in nota.runs:
        r.font.size = Pt(10)
    return nota


def main(src, dst):
    d = docx.Document(src)
    P = d.paragraphs
    borrar = []
    for i, (inicio, nuevos) in CAMBIOS.items():
        actual = P[i].text.strip()
        assert actual.startswith(inicio), (i, inicio, actual[:60])
        if not nuevos:
            borrar.append(i)
            continue
        R.reemplazar(P[i], nuevos[0])
        ancla = P[i]
        for texto in nuevos[1:]:
            nuevo = deepcopy(P[i]._p)
            ancla._p.addnext(nuevo)
            ancla = Paragraph(nuevo, P[i]._parent)
            R.reemplazar(ancla, texto)
    for a, b in ELIMINAR_RANGOS:
        for i in range(a, b + 1):
            if P[i].text.strip() and not P[i].style.name.startswith("Heading"):
                borrar.append(i)
    insertar_tabla(d, P[TABLA_TRAS], P[TABLA_TRAS])
    for i in sorted(set(borrar), reverse=True):
        P[i]._p.getparent().remove(P[i]._p)
    d.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
