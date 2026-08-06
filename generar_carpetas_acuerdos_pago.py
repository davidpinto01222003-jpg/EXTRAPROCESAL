# -*- coding: utf-8 -*-
"""
Genera el documento Word para marcar las carpetas fisicas de ACUERDOS DE PAGO.

Lee el Excel de acuerdos de pago y arma un Word con una hoja por periodo:
el MES, el ANO y la tabla de cuentas de energia relacionadas en ese mes.
El periodo se toma de la columna "ANO" del Excel, que viene como celdas
combinadas que abarcan todas las filas del mes correspondiente.

Uso:
    python generar_carpetas_acuerdos_pago.py
    python generar_carpetas_acuerdos_pago.py ACUERDOS_DE_PAGO.xlsx SALIDA.docx
"""

import os
import sys
from datetime import datetime

import openpyxl
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

EXCEL_POR_DEFECTO = "ACUERDOS_DE_PAGO.xlsx"
WORD_POR_DEFECTO = "ACUERDOS_DE_PAGO_CARPETAS.docx"

FILA_ENCABEZADO = 2      # fila con los titulos (N°, CUENTA, ...)
COL_N = 1
COL_CUENTA = 2
COL_DEMANDADO = 3
COL_MUNICIPIO = 4
COL_SERVICIO = 5
COL_PERIODO = 6          # columna F: "ANO"

MESES = {
    1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL", 5: "MAYO", 6: "JUNIO",
    7: "JULIO", 8: "AGOSTO", 9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE",
    12: "DICIEMBRE",
}

AZUL = RGBColor(0x1F, 0x38, 0x64)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
NEGRO = RGBColor(0x00, 0x00, 0x00)
GRIS_TEXTO = RGBColor(0x44, 0x44, 0x44)

FONDO_AZUL = "1F3864"
FONDO_AZUL_CLARO = "D9E2F3"
FONDO_GRIS = "F2F2F2"
BORDE_AZUL = "8EA9DB"

FUENTE = "Arial"

# Anchos de columna de la tabla de cuentas (pulgadas), suman 7.5"
ANCHOS = [0.6, 1.2, 3.1, 1.5, 1.1]
TITULOS = ["N°", "CUENTA", "DEMANDADO", "MUNICIPIO", "SERVICIO"]


# --------------------------------------------------------------------------
# Utilidades de formato (python-docx no expone sombreado ni bordes de tabla)
# --------------------------------------------------------------------------

def sombrear(celda, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    celda._tc.get_or_add_tcPr().append(shd)


def bordes_tabla(tabla, color_hex, grosor=6):
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(grosor))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color_hex)
        borders.append(el)
    tabla._tbl.tblPr.append(borders)


def repetir_encabezado(fila):
    trPr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def escribir(celda, texto, tamano=10, negrita=False, color=NEGRO,
             alineacion=WD_ALIGN_PARAGRAPH.LEFT, espaciado=None):
    parrafo = celda.paragraphs[0]
    parrafo.alignment = alineacion
    parrafo.paragraph_format.space_before = Pt(1)
    parrafo.paragraph_format.space_after = Pt(1)
    run = parrafo.add_run(texto)
    run.font.name = FUENTE
    run.font.size = Pt(tamano)
    run.font.bold = negrita
    run.font.color.rgb = color
    if espaciado:
        rPr = run._r.get_or_add_rPr()
        el = OxmlElement("w:spacing")
        el.set(qn("w:val"), str(espaciado))
        rPr.append(el)
    celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return run


def parrafo(doc, texto, tamano=11, negrita=False, color=NEGRO,
            alineacion=WD_ALIGN_PARAGRAPH.CENTER, antes=0, despues=6):
    p = doc.add_paragraph()
    p.alignment = alineacion
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    run = p.add_run(texto)
    run.font.name = FUENTE
    run.font.size = Pt(tamano)
    run.font.bold = negrita
    run.font.color.rgb = color
    return p


# --------------------------------------------------------------------------
# Lectura del Excel
# --------------------------------------------------------------------------

def limpiar(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return " ".join(str(valor).split())


def mapa_periodos(ws):
    """Devuelve {fila: valor de periodo} expandiendo las celdas combinadas."""
    periodos = {}
    for rango in ws.merged_cells.ranges:
        if rango.min_col != COL_PERIODO:
            continue
        valor = ws.cell(rango.min_row, COL_PERIODO).value
        for fila in range(rango.min_row, rango.max_row + 1):
            periodos[fila] = valor
    for fila in range(FILA_ENCABEZADO + 1, ws.max_row + 1):
        valor = ws.cell(fila, COL_PERIODO).value
        if valor is not None and fila not in periodos:
            periodos[fila] = valor
    return periodos


def etiqueta_periodo(valor):
    """Convierte el valor de la columna ANO en (mes, anio)."""
    if isinstance(valor, datetime):
        return MESES[valor.month], str(valor.year)
    texto = limpiar(valor).upper()
    digitos = "".join(c for c in texto if c.isdigit())
    return "", (digitos if digitos else texto)


def leer_periodos(ruta_excel):
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    ws = wb.worksheets[0]
    periodos = mapa_periodos(ws)

    grupos = []
    indice = {}
    ultimo = None
    for fila in range(FILA_ENCABEZADO + 1, ws.max_row + 1):
        cuenta = ws.cell(fila, COL_CUENTA).value
        demandado = ws.cell(fila, COL_DEMANDADO).value
        if cuenta is None and demandado is None:
            continue

        valor = periodos.get(fila, ultimo)
        if valor is None:
            continue
        ultimo = valor

        mes, anio = etiqueta_periodo(valor)
        clave = (mes, anio)
        if clave not in indice:
            indice[clave] = len(grupos)
            grupos.append({"mes": mes, "anio": anio, "cuentas": []})
        grupos[indice[clave]]["cuentas"].append({
            "n": limpiar(ws.cell(fila, COL_N).value),
            "cuenta": limpiar(cuenta),
            "demandado": limpiar(demandado),
            "municipio": limpiar(ws.cell(fila, COL_MUNICIPIO).value),
            "servicio": limpiar(ws.cell(fila, COL_SERVICIO).value),
        })

    return grupos


# --------------------------------------------------------------------------
# Armado del Word
# --------------------------------------------------------------------------

def bloque_periodo(doc, grupo):
    """Recuadro grande con MES y ANO, para recortar y pegar en la carpeta."""
    tabla = doc.add_table(rows=1, cols=1)
    tabla.autofit = False
    celda = tabla.rows[0].cells[0]
    celda.width = Inches(7.5)
    sombrear(celda, FONDO_AZUL_CLARO)
    bordes_tabla(tabla, FONDO_AZUL, grosor=18)

    celda.paragraphs[0]._p.getparent().remove(celda.paragraphs[0]._p)

    lineas = [
        ("ACUERDOS DE PAGO", 14, AZUL, 6, 3),
        (grupo["mes"] or "SIN MES REGISTRADO", 36, AZUL, 0, 2),
        (grupo["anio"], 28, NEGRO, 0, 6),
    ]
    for texto, tamano, color, antes, despues in lineas:
        p = celda.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(despues)
        run = p.add_run(texto)
        run.font.name = FUENTE
        run.font.size = Pt(tamano)
        run.font.bold = True
        run.font.color.rgb = color
    return tabla


def tabla_cuentas(doc, cuentas):
    tabla = doc.add_table(rows=1, cols=len(TITULOS))
    tabla.autofit = False
    bordes_tabla(tabla, BORDE_AZUL, grosor=6)

    encabezado = tabla.rows[0]
    repetir_encabezado(encabezado)
    for i, titulo in enumerate(TITULOS):
        celda = encabezado.cells[i]
        celda.width = Inches(ANCHOS[i])
        sombrear(celda, FONDO_AZUL)
        escribir(celda, titulo, tamano=10, negrita=True, color=BLANCO,
                 alineacion=WD_ALIGN_PARAGRAPH.LEFT if i == 2 else WD_ALIGN_PARAGRAPH.CENTER)

    for j, item in enumerate(cuentas):
        fila = tabla.add_row()
        fondo = FONDO_GRIS if j % 2 else None
        valores = [
            (item["n"], 10, False, WD_ALIGN_PARAGRAPH.CENTER),
            (item["cuenta"], 11, True, WD_ALIGN_PARAGRAPH.CENTER),
            (item["demandado"], 10, False, WD_ALIGN_PARAGRAPH.LEFT),
            (item["municipio"], 10, False, WD_ALIGN_PARAGRAPH.CENTER),
            (item["servicio"], 9, False, WD_ALIGN_PARAGRAPH.CENTER),
        ]
        for i, (texto, tamano, negrita, alineacion) in enumerate(valores):
            celda = fila.cells[i]
            celda.width = Inches(ANCHOS[i])
            if fondo:
                sombrear(celda, fondo)
            escribir(celda, texto, tamano=tamano, negrita=negrita, alineacion=alineacion)
    return tabla


def portada(doc, grupos):
    total = sum(len(g["cuentas"]) for g in grupos)
    parrafo(doc, "ACUERDOS DE PAGO", tamano=28, negrita=True, color=AZUL, antes=24, despues=4)
    parrafo(doc, "MARCACIÓN DE CARPETAS FÍSICAS POR MES Y AÑO",
            tamano=13, negrita=True, despues=16)
    parrafo(doc, f"{len(grupos)} carpetas · {total} cuentas de energía",
            tamano=12, color=GRIS_TEXTO, despues=18)

    anchos = [0.8, 2.9, 1.6, 2.2]
    titulos = ["N°", "MES", "AÑO", "CUENTAS DE ENERGÍA"]
    tabla = doc.add_table(rows=1, cols=len(titulos))
    tabla.autofit = False
    bordes_tabla(tabla, BORDE_AZUL, grosor=6)

    encabezado = tabla.rows[0]
    repetir_encabezado(encabezado)
    for i, titulo in enumerate(titulos):
        celda = encabezado.cells[i]
        celda.width = Inches(anchos[i])
        sombrear(celda, FONDO_AZUL)
        escribir(celda, titulo, tamano=10, negrita=True, color=BLANCO,
                 alineacion=WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER)

    for j, grupo in enumerate(grupos):
        fila = tabla.add_row()
        fondo = FONDO_GRIS if j % 2 else None
        valores = [
            (str(j + 1), WD_ALIGN_PARAGRAPH.CENTER, False),
            (grupo["mes"] or "SIN MES", WD_ALIGN_PARAGRAPH.LEFT, True),
            (grupo["anio"], WD_ALIGN_PARAGRAPH.CENTER, False),
            (str(len(grupo["cuentas"])), WD_ALIGN_PARAGRAPH.CENTER, False),
        ]
        for i, (texto, alineacion, negrita) in enumerate(valores):
            celda = fila.cells[i]
            celda.width = Inches(anchos[i])
            if fondo:
                sombrear(celda, fondo)
            escribir(celda, texto, tamano=11, negrita=negrita, alineacion=alineacion)


def salto_de_pagina(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def construir_word(grupos, ruta_docx):
    doc = Document()
    seccion = doc.sections[0]
    seccion.page_width = Inches(8.5)
    seccion.page_height = Inches(11)
    for lado in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(seccion, lado, Inches(0.5))

    estilo = doc.styles["Normal"]
    estilo.font.name = FUENTE
    estilo.font.size = Pt(10)

    portada(doc, grupos)

    for grupo in grupos:
        salto_de_pagina(doc)
        bloque_periodo(doc, grupo)
        parrafo(doc, f"CUENTAS DE ENERGÍA RELACIONADAS: {len(grupo['cuentas'])}",
                tamano=12, negrita=True, color=AZUL, antes=12, despues=8)
        tabla_cuentas(doc, grupo["cuentas"])

    doc.save(ruta_docx)


def main():
    ruta_excel = sys.argv[1] if len(sys.argv) > 1 else EXCEL_POR_DEFECTO
    ruta_docx = sys.argv[2] if len(sys.argv) > 2 else WORD_POR_DEFECTO

    if not os.path.exists(ruta_excel):
        print(f"[ERROR] No se encontro el Excel: {ruta_excel}")
        print("        Copie el archivo junto al script o pase la ruta como parametro.")
        sys.exit(1)

    grupos = leer_periodos(ruta_excel)
    if not grupos:
        print("[ERROR] No se encontraron periodos en la columna ANO del Excel.")
        sys.exit(1)

    construir_word(grupos, ruta_docx)

    total = sum(len(g["cuentas"]) for g in grupos)
    print(f"[OK] {ruta_docx}")
    print(f"     {len(grupos)} carpetas (hojas) | {total} cuentas de energia")
    for g in grupos:
        print(f"     - {(g['mes'] or 'SIN MES'):<12} {g['anio']}  {len(g['cuentas']):>3} cuentas")


if __name__ == "__main__":
    main()
