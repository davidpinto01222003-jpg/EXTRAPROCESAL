"""Construye la base de datos unificada a partir del COMPENDIO y de los libros de espectros."""
import openpyxl, json, math
import pandas as pd
import numpy as np

U = "/root/.claude/uploads/85bb4587-828c-5654-b9e1-8dcea31c6b7f/"
COMP = U + "f48aa56f-COMPENDIO_BASE_DE_DATOS.xlsx"

ERR = ("#DIV/0!", "#VALUE!", "#N/A", "ND", None, "")

def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return np.nan

# ---------- 1. Compendio ----------
wb = openpyxl.load_workbook(COMP, data_only=True)
blocks = []   # (hoja, anio, fila_encabezado)
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if isinstance(r[0], str) and r[0].startswith("FECHA"):
            blocks.append((ws.title, r[0].split("-")[1], i, rows))

recs = []
for hoja, anio, i, rows in blocks:
    labels = rows[i]
    codes = rows[i + 1]
    temp = rows[i + 2]
    sal = rows[i + 3]
    tur = rows[i + 4]
    br1 = rows[i + 5]
    br2 = rows[i + 6]
    for c in range(2, len(codes)):
        code = codes[c]
        if not isinstance(code, str) or not code.startswith("SC-"):
            continue
        recs.append(dict(
            epoca="Seca" if hoja == "SECA" else "Humeda",
            anio=int(anio),
            campana=f"{'Seca' if hoja=='SECA' else 'Humeda'} {anio}",
            etiqueta=labels[c],
            estacion=code,
            temperatura=num(temp[c]),
            salinidad=num(sal[c]),
            turbidez=num(tur[c]),
            br440_compendio=num(br1[c]),
            br443_compendio=num(br2[c]),
            br440_flag=(br1[c] if isinstance(br1[c], str) else ""),
        ))
df = pd.DataFrame(recs)

# normalizacion de codigos SC-21-225 -> SC-21-0225
def norm(code):
    a, b, c = code.split("-")
    return f"{a}-{b}-{int(c):04d}"
df["estacion"] = df["estacion"].map(norm)

# ---------- 2. Espectros ----------
spec_files = {
    "Seca 2021": "f13fb917-Espectros_apadaphyCDOM_Sol001321.xlsx",
    "Seca 2022": "a38edcf6-Espectros_apadaphyCDOM_Sol000822_1.xlsx",
    "Humeda 2022": "fba8bf6e-Espectros_apadaphyCDOM_Sol003422_1.xlsx",
    "Humeda 2023": "952b6f01-Espectros_apadaphyCDOM_Sol002423_1.xlsx",
}

def leer_hoja(ws, ncols_max=None):
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    wl = {r[0]: r for r in rows[1:] if isinstance(r[0], (int, float))}
    return hdr, wl

spec = []
for camp, f in spec_files.items():
    wbs = openpyxl.load_workbook(U + f, read_only=True, data_only=True)
    codes_ap = [c for c in next(wbs["ap"].iter_rows(max_row=1, values_only=True))[1:] if c]
    hdr, wl = leer_hoja(wbs["aphy"])
    # primer bloque: columnas 1..len(codes_ap)
    for j, code in enumerate(codes_ap, start=1):
        rec = dict(campana=camp, col=j, header=hdr[j])
        for lam in (440, 443, 675, 676, 677):
            r = wl.get(lam)
            rec[f"aphy{lam}"] = num(r[j]) if r else np.nan
        # ad y ap para control de calidad
        for hoja, pref in (("ap", "ap"), ("ad", "ad")):
            h2, w2 = leer_hoja(wbs[hoja])
            for lam in (440, 676):
                r = w2.get(lam)
                rec[f"{pref}{lam}"] = num(r[j]) if r and j < len(r) else np.nan
        spec.append(rec)
    wbs.close()
sp = pd.DataFrame(spec)

# mapeo columna -> estacion usando el orden del compendio por campana
orden = {}
for camp in spec_files:
    sub = df[df.campana == camp]
    orden[camp] = list(sub.estacion)
sp["estacion"] = [orden[r.campana][r.col - 1] if r.col - 1 < len(orden[r.campana]) else None
                  for r in sp.itertuples()]

sp["br440_espectro"] = sp.aphy440 / sp.aphy676
sp["br443_espectro"] = sp.aphy443 / sp.aphy677

out = df.merge(sp[["campana", "estacion", "aphy440", "aphy443", "aphy675", "aphy676", "aphy677",
                   "ap440", "ap676", "ad440", "ad676",
                   "br440_espectro", "br443_espectro"]],
               on=["campana", "estacion"], how="left")
out.to_csv("base_unificada.csv", index=False)
print(out.groupby("campana").size())
print(out.head(8).to_string())
print("\nfilas:", len(out), " con espectro:", out.br440_espectro.notna().sum())
