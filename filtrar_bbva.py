"""
Toma el Excel que ya genero depurar_fng_leasing.py (con FNG y leasing de
CUALQUIER entidad) y lo depura para dejar SOLO las garantias FNG y los
leasing a favor de BBVA -- sin volver a leer el Drive.

Como depurar_fng_leasing.py guarda el texto completo de cada documento
leido (depurar_fng_leasing_cache.jsonl), aqui solo se vuelven a analizar,
con ese texto guardado, los documentos de las hojas FNG_DETALLE y
LEASING_DETALLE, aplicando el criterio de ENTIDAD_OBJETIVO de
depurar_fng_leasing.py (BBVA / Banco Bilbao Vizcaya cerca de la mencion,
o el documento en una carpeta / con un nombre que diga BBVA). Tarda
segundos.

Genera un Excel nuevo con el mismo formato, terminado en _BBVA.xlsx. El
Excel original no se toca. Los procesos que solo tenian FNG/leasing de
otras entidades quedan en NO, con una nota en OBSERVACIONES.

Uso:
    python filtrar_bbva.py                 (usa el depuracion_fng_leasing_*.xlsx mas reciente)
    python filtrar_bbva.py "C:\\ruta\\depuracion_fng_leasing_2026-10-02.xlsx"
"""

import logging
import os
import re
import sys
from pathlib import Path

import openpyxl

import depurar_fng_leasing as d

ARCHIVO_LOG = os.path.join(d.DIRECTORIO, "filtrar_bbva.log")


def buscar_excel_resultado():
    if len(sys.argv) > 1:
        return sys.argv[1].strip().strip('"')
    candidatos = [
        p for p in Path(d.DIRECTORIO).glob("depuracion_fng_leasing_*.xlsx")
        if not p.stem.upper().endswith("_BBVA") and not p.name.startswith("~$")
    ]
    if not candidatos:
        raise RuntimeError(
            "No encontre ningun depuracion_fng_leasing_*.xlsx en esta carpeta. "
            "Primero deja terminar depurar_fng_leasing.py (EJECUTAR.bat)."
        )
    return str(max(candidatos, key=lambda p: p.stat().st_mtime))


def id_de_enlace(url):
    """ID del archivo de Drive a partir de su enlace (o la ruta local tal cual)."""
    if not url:
        return None
    m = re.search(r"/d/([A-Za-z0-9_\-]+)", url) or re.search(r"[?&]id=([A-Za-z0-9_\-]+)", url)
    if m:
        return m.group(1)
    return url if not url.startswith("http") else None


def _filas_con_enlaces(hoja, col_enlace):
    """[(dict encabezado->valor, enlace)] de una hoja de detalle."""
    filas = list(hoja.iter_rows(min_row=1))
    if not filas:
        return []
    encabezados = [c.value for c in filas[0]]
    salida = []
    for fila in filas[1:]:
        valores = {encabezados[i]: c.value for i, c in enumerate(fila) if i < len(encabezados)}
        if not valores.get("CONCURSADO"):
            continue
        celda = fila[col_enlace]
        enlace = celda.hyperlink.target if celda.hyperlink else (celda.value if celda.value != "Abrir" else None)
        salida.append((valores, enlace))
    return salida


def _archivo(valores, enlace):
    ruta_completa = str(valores.get("RUTA EN LA CARPETA") or "")
    carpeta, _, ruta = ruta_completa.partition("/")
    return {"nombre": valores.get("DOCUMENTO") or "", "carpeta": carpeta, "ruta": ruta or ruta_completa,
            "enlace": enlace or "", "id": id_de_enlace(enlace)}


def main():
    logging.basicConfig(
        level=logging.INFO, format="%(message)s",
        handlers=[logging.FileHandler(ARCHIVO_LOG, encoding="utf-8"), logging.StreamHandler()],
    )
    for nombre in ("pypdf", "PyPDF2"):
        logging.getLogger(nombre).setLevel(logging.ERROR)
    if d.RE_ENTIDAD_OBJETIVO is None:
        raise RuntimeError("ENTIDAD_OBJETIVO esta vacia en depurar_fng_leasing.py: no hay entidad por la cual filtrar.")

    ruta = buscar_excel_resultado()
    salida = str(Path(ruta).with_name(Path(ruta).stem + "_BBVA.xlsx"))
    logging.info("Excel a depurar: %s", ruta)
    libro = openpyxl.load_workbook(ruta)  # no read_only: hacen falta los hipervinculos

    logging.info("Cargando el texto ya leido de los documentos (cache)...")
    cache = d.Cache(d.ARCHIVO_CACHE)
    logging.info("   %d documentos en el cache.", len(cache.datos))

    # --- PROCESOS: listado original + columnas de la corrida ---
    hp = libro["PROCESOS"]
    filas_p = list(hp.iter_rows(min_row=1))
    enc = [c.value for c in filas_p[0]]
    n_orig = enc.index("¿EN DRIVE?")
    encabezados = [str(e) for e in enc[:n_orig]]
    col = {e: i for i, e in enumerate(enc)}
    idx_conc = next(
        (i for i, e in enumerate(encabezados) if d._encabezado_norm(e) in d.ENCABEZADOS_CONCURSADO), 0
    )
    idx_tipo = next((i for i, e in enumerate(encabezados) if d._encabezado_norm(e) in d.ENCABEZADOS_TIPO), None)

    no_encontrados = {}
    if "NO_ENCONTRADOS_EN_DRIVE" in libro.sheetnames:
        for valores, enlace in _filas_con_enlaces(libro["NO_ENCONTRADOS_EN_DRIVE"], 6):
            no_encontrados.setdefault(str(valores["CONCURSADO"]), (valores, enlace))
    sin_leer_por_proceso = {}
    if "REVISAR_A_MANO" in libro.sheetnames:
        for valores, enlace in _filas_con_enlaces(libro["REVISAR_A_MANO"], 3):
            sin_leer_por_proceso.setdefault(str(valores["CONCURSADO"]), []).append(
                (_archivo(valores, enlace), valores.get("MOTIVO") or "")
            )

    resultados = []
    for fila in filas_p[1:]:
        valores = [c.value for c in fila]
        if not any(v not in (None, "") for v in valores[:n_orig]):
            continue
        nombre = str(valores[idx_conc] or "").strip()
        nombres_carpetas = [x for x in str(valores[col["CARPETA(S) EN DRIVE"]] or "").split(" | ") if x]
        celda_enlace = fila[col["ENLACE CARPETA"]]
        enlace_carpeta = celda_enlace.hyperlink.target if celda_enlace.hyperlink else ""
        obs = str(valores[col["OBSERVACIONES"]] or "")
        res = {
            "fila": valores[:n_orig], "nombre": nombre,
            "tipo": str(valores[idx_tipo] or "").strip() if idx_tipo is not None else "",
            "nit": "", "expedientes": [],
            "carpetas": [{"nombre": n, "enlace": enlace_carpeta if i == 0 else ""}
                         for i, n in enumerate(nombres_carpetas)],
            "alternativas": [], "fng": [], "leasing": [],
            "sin_leer": sin_leer_por_proceso.get(nombre, []),
            "total_archivos": valores[col["ARCHIVOS EN CARPETA"]] or "",
            "observaciones": obs, "otros": {"fng": 0, "leasing": 0}, "sin_verificar": 0,
        }
        if nombre in no_encontrados:
            v, enl = no_encontrados[nombre]
            res["nit"] = v.get("NIT") or ""
            res["expedientes"] = [x for x in str(v.get("EXPEDIENTE") or "").split(", ") if x]
            if v.get("CARPETA MAS PARECIDA"):
                res["alternativas"] = [(float(v.get("SIMILITUD") or 0), {"nombre": v["CARPETA MAS PARECIDA"],
                                                                         "enlace": enl or ""})]
        resultados.append(res)

    por_nombre = {}
    for r in resultados:
        por_nombre.setdefault(r["nombre"], []).append(r)

    # --- Re-analiza cada documento de las hojas de detalle ---
    revisados = {"fng": 0, "leasing": 0}
    for hoja, clave, analizar in (("FNG_DETALLE", "fng", d.analizar_fng),
                                  ("LEASING_DETALLE", "leasing", d.analizar_leasing)):
        if hoja not in libro.sheetnames:
            continue
        for valores, enlace in _filas_con_enlaces(libro[hoja], 3):
            procesos = [r for r in por_nombre.get(str(valores["CONCURSADO"]).strip(), []) if r["carpetas"]]
            if not procesos:
                continue
            a = _archivo(valores, enlace)
            nombre_y_ruta = f"{a['carpeta']}/{a['ruta']}"
            entrada = cache.datos.get(a["id"]) if a["id"] else None
            texto = entrada["texto"] if entrada else ""
            revisados[clave] += 1
            h = analizar(texto, nombre_y_ruta)
            for r in procesos:
                if h and h.get("objetivo"):
                    r[clave].append((a, h))
                elif entrada is None and not d.RE_ENTIDAD_OBJETIVO.search(d.normalizar(nombre_y_ruta)):
                    r["sin_verificar"] += 1
                    r["sin_leer"].append((a, f"Menciona {clave.upper()} pero no se pudo verificar si es de "
                                             f"{d.ETIQUETA_ENTIDAD} (texto no guardado): revisar a mano"))
                else:
                    r["otros"][clave] += 1

    for r in resultados:
        if not r["carpetas"]:
            continue
        r["resumen"] = d.resumir_proceso(r["fng"], r["leasing"])
        notas = [r["observaciones"]] if r["observaciones"] else []
        if r["otros"]["fng"] or r["otros"]["leasing"]:
            partes = []
            if r["otros"]["fng"]:
                partes.append(f"FNG en {r['otros']['fng']} documento(s)")
            if r["otros"]["leasing"]:
                partes.append(f"leasing en {r['otros']['leasing']} documento(s)")
            notas.append(f"Menciona {' y '.join(partes)} SIN {d.ETIQUETA_ENTIDAD} cerca "
                         "(otra entidad; no se cuentan)")
        if r["sin_verificar"]:
            notas.append(f"{r['sin_verificar']} documento(s) sin verificar (ver REVISAR_A_MANO)")
        r["observaciones"] = ". ".join(notas)

    total_filas = None
    if "RESUMEN" in libro.sheetnames:
        for fila in libro["RESUMEN"].iter_rows(values_only=True):
            if fila and fila[0] == "Filas del listado original":
                total_filas = fila[1]
    contadores = d.calcular_contadores(resultados, total_filas)
    d.escribir_excel(salida, encabezados, resultados, contadores)

    logging.info("")
    logging.info("Documentos re-analizados: %d con FNG, %d con leasing.", revisados["fng"], revisados["leasing"])
    for etiqueta, valor in contadores:
        logging.info("%-58s %s", etiqueta, valor)
    logging.info("")
    logging.info("Listo. Resultado solo %s en: %s", d.ETIQUETA_ENTIDAD, salida)
    if sys.platform.startswith("win"):
        try:
            os.startfile(salida)
        except OSError:
            pass


if __name__ == "__main__":
    try:
        main()
    except Exception as e:  # noqa: BLE001
        logging.basicConfig(level=logging.INFO, format="%(message)s")
        logging.exception("ERROR: %s", e)
        sys.exit(1)
