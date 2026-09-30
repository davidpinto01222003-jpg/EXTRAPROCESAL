"""Simulador aproximado de Compilatio (IA + similitud) calibrado con un informe real.

No es Compilatio: aprende qué frases marcó Compilatio en un informe real y
estima cómo marcaría una versión nueva del mismo documento.

Uso:
    python simulador.py entrenar datos/original.docx datos/reporte_compilatio.pdf \
                                 datos/v4.docx datos/reporte_compilatio_v4.pdf
    python simulador.py evaluar "ruta/al/trabajo.docx" --html reporte.html
"""
import argparse
import html
import pickle
import re
import statistics
import sys
from pathlib import Path

import docx
import numpy as np
import pdfplumber
from sklearn.ensemble import HistGradientBoostingClassifier
from sklearn.metrics import roc_auc_score
from sklearn.model_selection import GroupKFold, cross_val_predict

AQUI = Path(__file__).parent
MODELO = AQUI / "modelo.pkl"
sys.path.insert(0, str(AQUI))

# Detectores de Hugging Face que se añaden con `entrenar --hf` (ver detectores_hf.py)
HF_USADOS = ["xlmr_es", "autext", "binoculars", "ppl"]

AZUL_IA = "(0.0, 0.749, 1.0)"
ROSA_IDIOMA = "(1.0, 0.3569, 0.7059)"
BLANCO = "(1.0, 1.0, 1.0)"
GRIS_TABLA = "(0.1169, 0.1608, 0.222)"

CONECTORES = [
    "conviene", "esto es", "de ahí", "en consecuencia", "por su parte", "asimismo",
    "en ese sentido", "cabe", "resulta", "se trata de", "no solo", "sino también",
    "por lo anterior", "en efecto", "a su turno", "en este sentido", "ahora bien",
    "de modo que", "de manera que", "en la medida en que", "por tanto", "así,",
    "en primer término", "la primera", "la segunda", "la tercera", "en síntesis",
    "en suma", "finalmente", "además", "sin embargo", "por el contrario",
    "en particular", "es decir", "a juicio de", "en el plano", "frente a",
]
GENERICAS = [
    "fundamental", "clave", "relevante", "central", "estructural", "integral",
    "sostenible", "efectiva", "eficacia", "articulación", "articulada", "abordar",
    "dimensión", "fenómeno", "ámbito", "marco", "garantizar", "coordinación",
    "consolidado", "significativ", "robust", "precisamente", "evidenci",
]
HUMANAS = [
    "nosotros", "creemos", "revisamos", "encontramos", "nuestro", "nuestra",
    "a nuestro juicio", "es cierto", "eso sí", "o sea", "pero ", "y justo",
    "?", "¿", "pasa", "sale", "queda", "casi",
]


def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def tokens(s):
    return re.findall(r"\w+", s.lower())


def shingles(s, n=3):
    t = tokens(s)
    return {tuple(t[i:i + n]) for i in range(max(0, len(t) - n + 1))}


def oraciones(parrafo):
    partes = re.split(r"(?<=[.;?!])\s+(?=[A-ZÁÉÍÓÚÑ¿“(])", parrafo)
    return [p for p in (norm(x) for x in partes) if p]


def leer_docx(ruta):
    d = docx.Document(ruta)
    pars = [norm(p.text) for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                pars.append(norm(c.text))
    return [p for p in pars if p]


def cuerpo(pars):
    """Párrafos que cuentan para IA: prosa (>=25 palabras), antes de la bibliografía."""
    fuera = False
    out = []
    for i, p in enumerate(pars):
        if re.fullmatch(r"(Referencias|Bibliografía|REFERENCIAS)", p):
            fuera = True
        if not fuera and len(p.split()) >= 25:
            out.append((i, p))
    return out


# ---------------------------------------------------------------- etiquetas

def marcas_pdf(pdf):
    ia, sim, idioma = [], [], []
    with pdfplumber.open(pdf) as doc:
        for pg in doc.pages[5:]:
            words = pg.extract_words()
            for r in pg.rects:
                col = str(r.get("non_stroking_color"))
                if r["x0"] < 60 or col in (BLANCO, GRIS_TABLA):
                    continue
                if col in (AZUL_IA, ROSA_IDIOMA):
                    ws = [w for w in words if 0 <= r["top"] - w["bottom"] <= 6
                          and w["x0"] >= r["x0"] - 2 and w["x1"] <= r["x1"] + 2]
                    (ia if col == AZUL_IA else idioma).append(" ".join(w["text"] for w in ws))
                elif r["bottom"] - r["top"] >= 5:
                    ws = [w for w in words if w["top"] >= r["top"] - 2 and w["bottom"] <= r["bottom"] + 2
                          and w["x0"] >= r["x0"] - 2 and w["x1"] <= r["x1"] + 2]
                    sim.append(" ".join(w["text"] for w in ws))
    return [x for x in ia if x], [x for x in sim if len(x.split()) >= 4], [x for x in idioma if x]


def porcentajes_pdf(pdf):
    """Lee del resumen del informe los porcentajes de IA, similitud e idiomas."""
    with pdfplumber.open(pdf) as doc:
        txt = doc.pages[0].extract_text() or ""
    def pct(rotulo):
        m = re.search(rotulo + r"\s+<?(\d+)\s*%", txt)
        return float(m.group(1)) if m else 0.0
    return {"ia": pct("Detección de IA"), "sim": pct("Similitudes"), "idioma": pct("Idiomas no reconocidos")}


def etiquetar(pc, lineas_ia):
    """Marca cada oración del cuerpo con 1 si Compilatio subrayó en azul al menos la mitad."""
    sh_ia = set()
    for l in lineas_ia:
        sh_ia |= shingles(l)
    pares, y = [], []
    for i, p in pc:
        for s in oraciones(p):
            sh = shingles(s)
            if not sh:
                continue
            pares.append((i, s, p))
            y.append(1 if len(sh & sh_ia) / len(sh) >= 0.5 else 0)
    return pares, np.array(y)


# ---------------------------------------------------------------- rasgos

NOMBRES_RASGOS = [
    "palabras", "log palabras", "comas/palabra", "; y :", "citas (año)", "conectores", "genéricas",
    "marcas humanas", "variedad léxica", "largo medio de palabra", "palabras >=10 letras", "se + verbo",
    "y/o", "variación largo oraciones del párrafo", "palabras del párrafo", "tiene dígitos",
    "posición en el párrafo", "oraciones del párrafo", "largo relativo", "posición en el documento",
    "empieza con conector", "dos puntos", "paréntesis", "comillas", "primera persona plural",
    "-ción/-miento/-idad", "es/son/fue/era", "rango de largos del párrafo", "termina en dos puntos",
]


def rasgos(s, p, pos_doc=0.5):
    ts = tokens(s)
    n = max(1, len(ts))
    low = s.lower()
    os_ = oraciones(p)
    largos_p = [len(tokens(o)) for o in os_] or [n]
    cv = statistics.pstdev(largos_p) / (statistics.mean(largos_p) or 1)
    k = os_.index(s) if s in os_ else 0
    return [
        n,
        np.log1p(n),
        s.count(",") / n,
        s.count(";") + s.count(":"),
        len(re.findall(r"\(\w[^)]*\d{4}[a-z]?\)", s)),
        sum(low.count(c) for c in CONECTORES),
        sum(low.count(g) for g in GENERICAS) / n * 10,
        sum(low.count(h) for h in HUMANAS),
        len(set(ts)) / n,
        np.mean([len(t) for t in ts]) if ts else 0,
        sum(len(t) >= 10 for t in ts) / n,
        len(re.findall(r"\bse \w+", low)) / n * 10,
        len(re.findall(r"\b(y|o)\b", low)) / n * 10,
        cv,
        len(p.split()),
        1 if re.search(r"\d", s) else 0,
        k / (len(os_) - 1) if len(os_) > 1 else 0,
        len(os_),
        n / statistics.mean(largos_p),
        pos_doc,
        1 if re.match(r"^(así|además|por|en|de|con|sin embargo|ahora|pero|y)\b", low) else 0,
        s.count(":"),
        s.count("("),
        1 if "“" in s or '"' in s else 0,
        len(re.findall(r"\b(nosotros|nuestr\w*|\w+amos)\b", low)),
        len(re.findall(r"\b\w+(ción|miento|idad)\b", low)) / n * 10,
        len(re.findall(r"\b(es|son|fue|era)\b", low)) / n * 10,
        max(largos_p) - min(largos_p),
        1 if s.endswith(":") else 0,
    ]


def matriz(pares, pc, hf=()):
    orden = {i: k / max(1, len(pc) - 1) for k, (i, _) in enumerate(pc)}
    X = np.array([rasgos(s, p, orden[i]) for i, s, p in pares])
    if hf:
        X = np.hstack([X, rasgos_hf([p for _, _, p in pares], hf)])
    return X


def rasgos_hf(parrafos, nombres):
    """Puntajes de los detectores de Hugging Face del párrafo de cada oración (opcional).

    Probados contra las marcas reales de la v4 no generalizan (AUC 0,57 entrenando
    con el original), por eso el modelo por defecto no los usa.
    """
    import detectores_hf as H
    unicos = sorted(set(parrafos))
    por_p = dict(zip(unicos, H.puntajes(unicos)))
    return np.array([[por_p[p][n] for n in nombres] for p in parrafos])


def nuevo_modelo(semilla=0):
    return HistGradientBoostingClassifier(max_depth=3, learning_rate=0.05, max_iter=200,
                                          class_weight="balanced", random_state=semilla)


# ---------------------------------------------------------------- comandos

def entrenar(pares_docx_pdf, hf=False):
    """Entrena con uno o más (trabajo .docx, informe .pdf de Compilatio), en orden cronológico."""
    nombres_hf = HF_USADOS if hf else []
    Xs, ys, grupos, memoria, informes = [], [], [], {}, []
    for k, (docx_path, pdf_path) in enumerate(pares_docx_pdf):
        pars = leer_docx(docx_path)
        pc = cuerpo(pars)
        ia, sim, _ = marcas_pdf(pdf_path)
        pares, y = etiquetar(pc, ia)
        Xs.append(matriz(pares, pc, nombres_hf))
        ys.append(y)
        grupos += [(k, i) for i, _, _ in pares]
        # memoria: la marca real de cada oración; el informe más reciente manda
        memoria.update({s: int(v) for (_, s, _), v in zip(pares, y)})
        palabras = np.array([len(s.split()) for _, s, _ in pares])
        total_doc = sum(len(p.split()) for p in pars)
        informes.append({"docx": Path(docx_path).name, "real": porcentajes_pdf(pdf_path),
                         "sim": fragmentos_presentes(sim, pars),
                         "ia_etiquetas": float((palabras * y).sum() / total_doc * 100),
                         "palabras": palabras, "total_doc": total_doc})
    X, y = np.vstack(Xs), np.concatenate(ys)
    # validación cruzada por párrafos: las oraciones de un mismo párrafo
    # comparten rasgos, y mezclarlas entre entrenamiento y prueba infla el AUC
    cv = GroupKFold(5).split(X, y, [hash(g) for g in grupos])
    prob_cv = cross_val_predict(nuevo_modelo(), X, y, cv=cv, method="predict_proba")[:, 1]
    auc = roc_auc_score(y, prob_cv)
    palabras = np.concatenate([r["palabras"] for r in informes])
    total = sum(r["total_doc"] for r in informes)
    real = (palabras * y).sum()
    # umbral con el que el modelo reproduce, en validación cruzada, el % de las etiquetas
    umbral = min(np.linspace(0.2, 0.9, 71), key=lambda u: abs((palabras * (prob_cv >= u)).sum() - real))
    fuera = None
    if len(informes) > 1:  # prueba honesta: entrenar con los informes anteriores y predecir el último
        n_ult = len(ys[-1])
        m = nuevo_modelo().fit(X[:-n_ult], y[:-n_ult])
        fuera = roc_auc_score(ys[-1], m.predict_proba(X[-n_ult:])[:, 1])
    modelo = nuevo_modelo().fit(X, y)
    ult = informes[-1]
    pickle.dump({"modelo": modelo, "umbral": float(umbral), "hf": nombres_hf, "memoria": memoria,
                 "sim": ult["sim"], "ref": ult["real"], "ref_ia_etiquetas": ult["ia_etiquetas"],
                 "auc_cv": float(auc), "auc_ultimo": fuera,
                 "informes": [{k: v for k, v in r.items() if k in ("docx", "real", "ia_etiquetas")} for r in informes]},
                open(MODELO, "wb"))
    for r in informes:
        print(f"{r['docx']}: Compilatio IA {r['real']['ia']:.0f} % | según etiquetas extraídas {r['ia_etiquetas']:.1f} %")
    print(f"Oraciones de entrenamiento: {len(y)} (marcadas IA: {y.sum()}) | memoria: {len(memoria)} oraciones")
    print(f"Detectores HF: {', '.join(nombres_hf) or 'no'} | umbral {umbral:.2f}")
    print(f"AUC validación cruzada agrupada: {auc:.3f}" + (f" | AUC prediciendo el último informe: {fuera:.3f}" if fuera else ""))


def fragmentos_presentes(fragmentos, pars):
    """Fragmentos de similitud que siguen en el texto fuera de comillas."""
    sin_comillas = re.sub(r"“[^”]*”|\"[^\"]*\"", " ", " ".join(pars))
    sh_doc = shingles(sin_comillas, 5)
    return [f for f in fragmentos if shingles(f, 5) and len(shingles(f, 5) & sh_doc) / len(shingles(f, 5)) >= 0.6]


def evaluar(docx_path, html_out=None, silencioso=False):
    M = pickle.load(open(MODELO, "rb"))
    pars = leer_docx(docx_path)
    total_doc = sum(len(p.split()) for p in pars)
    pc = cuerpo(pars)
    pares = [(i, s, p) for i, p in pc for s in oraciones(p)]
    probs = M["modelo"].predict_proba(matriz(pares, pc, M["hf"]))[:, 1]
    filas = []
    for (i, s, _), prob in zip(pares, probs):
        if s in M["memoria"]:  # oración sin cambios: se usa la marca real del último informe
            filas.append((i, s, float(M["memoria"][s]), bool(M["memoria"][s]), True))
        else:
            filas.append((i, s, prob, prob >= M["umbral"], False))
    ia_words = sum(len(s.split()) for _, s, _, f, _ in filas if f)
    ia_pct = ia_words / total_doc * 100 * (M["ref"]["ia"] / M["ref_ia_etiquetas"])

    # similitud: de los fragmentos que Compilatio encontró, cuántos siguen fuera de comillas
    quedan = fragmentos_presentes(M["sim"], pars)
    base = M["sim"]
    sim_pct = M["ref"]["sim"] * (sum(len(f.split()) for f in quedan) / max(1, sum(len(f.split()) for f in base)))
    idioma_pct = M["ref"]["idioma"]
    total = ia_pct + sim_pct + idioma_pct

    por_par = {}
    for i, s, prob, f, mem in filas:
        a = por_par.setdefault(i, [0, 0, []])
        a[0] += len(s.split()) * f
        a[1] += len(s.split())
        a[2].append((s, prob, f, mem))
    nuevas = [f for f in filas if not f[4]]
    if not silencioso:
        print(f"Palabras: {total_doc} | oraciones con marca conocida: {len(filas) - len(nuevas)}, nuevas: {len(nuevas)}")
        print(f"IA estimada ........... {ia_pct:5.1f} %")
        print(f"Similitud estimada .... {sim_pct:5.1f} %")
        print(f"Idiomas no reconocidos  {idioma_pct:5.1f} %  (se asume igual al último informe)")
        print(f"TOTAL estimado ........ {total:5.1f} %")
        print("\nPárrafos con más texto marcado como IA:")
        for i, (w, n, _) in sorted(por_par.items(), key=lambda kv: -kv[1][0])[:15]:
            if w:
                print(f"  [{i}] {w}/{n} palabras — {pars[i][:80]}…")

    if html_out:
        partes = []
        for i, (w, n, ss) in sorted(por_par.items()):
            frag = " ".join(
                f'<span class="{"ia" if f else ""}" title="{"marca real del informe" if mem else f"prob. IA {prob:.0%}"}">'
                f'{html.escape(s)}</span>' for s, prob, f, mem in ss)
            partes.append(f'<p><small>[{i}]</small> {frag}</p>')
        Path(html_out).write_text(PLANTILLA.format(
            nombre=html.escape(Path(docx_path).name), total=total, ia=ia_pct, sim=sim_pct,
            idioma=idioma_pct, palabras=total_doc, cuerpo="\n".join(partes)), encoding="utf8")
        if not silencioso:
            print(f"\nReporte HTML: {html_out}")
    return {"ia": ia_pct, "sim": sim_pct, "idioma": idioma_pct, "total": total, "por_par": por_par, "pars": pars}


PLANTILLA = """<!doctype html><html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Simulación Compilatio</title><style>
:root{{--bg:#fff;--fg:#1d1d1f;--mut:#666;--ia:#cdeefe;--line:#00a6e6;--card:#f5f6f8}}
@media (prefers-color-scheme:dark){{:root:not([data-theme="light"]){{--bg:#16181b;--fg:#e8e8ea;--mut:#9a9aa0;--ia:#0b3a4d;--line:#39c1f5;--card:#212428}}}}
body{{background:var(--bg);color:var(--fg);font:15px/1.6 system-ui,sans-serif;max-width:900px;margin:0 auto;padding:16px}}
.kpis{{display:flex;flex-wrap:wrap;gap:12px;margin:16px 0}}.k{{background:var(--card);border-radius:10px;padding:12px 16px;flex:1;min-width:140px}}
.k b{{display:block;font-size:28px}}.ia{{background:var(--ia);border-bottom:2px solid var(--line)}}small{{color:var(--mut)}}
</style></head><body><h1>Simulación tipo Compilatio</h1>
<p><small>{nombre} · {palabras} palabras · estimación aproximada calibrada con el informe real de Compilatio; no reemplaza el análisis oficial.</small></p>
<div class="kpis"><div class="k">Total estimado<b>{total:.1f} %</b></div><div class="k">IA<b>{ia:.1f} %</b></div>
<div class="k">Similitud<b>{sim:.1f} %</b></div><div class="k">Idiomas no reconocidos<b>{idioma:.1f} %</b></div></div>
<p><span class="ia">Texto resaltado</span> = frase marcada como IA: si no cambió desde el último informe, es la marca real de Compilatio; si es nueva, es la predicción del simulador (pasa el cursor para ver cuál).</p>
{cuerpo}</body></html>"""


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("entrenar", help="pares DOCX PDF en orden cronológico")
    e.add_argument("archivos", nargs="+")
    e.add_argument("--hf", action="store_true", help="añade los detectores de Hugging Face como rasgos")
    v = sub.add_parser("evaluar"); v.add_argument("docx"); v.add_argument("--html")
    a = ap.parse_args()
    if a.cmd == "entrenar":
        if len(a.archivos) % 2:
            ap.error("entrenar necesita pares: trabajo.docx informe.pdf [trabajo2.docx informe2.pdf ...]")
        entrenar(list(zip(a.archivos[::2], a.archivos[1::2])), hf=a.hf)
    else:
        evaluar(a.docx, a.html)
